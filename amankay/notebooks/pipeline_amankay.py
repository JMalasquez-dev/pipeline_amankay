# Databricks notebook source
# MAGIC %md
# MAGIC # Amankay - Pipeline de datos (fuente -> bronze -> silver -> gold -> integración)
# MAGIC Antes de ejecutar: crea el secret scope `amankay` con `aws_access_key`, `aws_secret_key` y `aws_session_token`
# MAGIC (ver README.md). El bucket se elige en el widget `bucket`.

# COMMAND ----------

# MAGIC %pip install -q boto3 shapely openpyxl xlsxwriter

# COMMAND ----------

import csv, functools, hashlib, html, io, json, os, re, threading, time, unicodedata
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta
from email.utils import parsedate_to_datetime
from getpass import getpass
from os.path import basename
from zoneinfo import ZoneInfo
import xml.etree.ElementTree as ET

import boto3
import pandas as pd
import requests
import shapely
from botocore.exceptions import ClientError
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from shapely.geometry import shape
from shapely.ops import unary_union
from shapely.strtree import STRtree

FUENTES = ["open_meteo", "igp", "indeci", "ingemmet", "reporte", "integracion"]
# Si el bucket YA tiene las demás fuentes (carpetas gold/ de cada una), basta con: FUENTES = ["reporte", "integracion"]
ETAPAS = ["fuente", "bronze", "silver", "gold"]
INICIO, FIN = date(2025, 1, 1), date(2025, 12, 31)
DEPARTAMENTOS = None  # solo Open-Meteo: None = todos, o por ejemplo ["AREQUIPA"]
HILOS = 4
TIPO_CRUCE = "left"  # IGP es la base: solo salen filas que existen en IGP; INDECI, INGEMMET y Open-Meteo solo agregan columnas

BUCKET_DEFECTO = "amankay-datalake-2026-v1"
REGION = "us-east-1"
CAPA_FUENTE = "fuente_de_datos"
PAUSA = 1.5
TZ = "America/Lima"
SEP = ";"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
UA = {"User-Agent": "Mozilla/5.0"}
UA_CHROME = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                           "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"}
LLAVES = ["fecha", "departamento", "provincia"]
ALIAS_ZONAS = {
    ("LIMA METROPOLITANA", "LIMA METROPOLITANA"): ("LIMA", "LIMA"),
    ("CALLAO", "PROV. CONST. DEL CALLAO"): ("CALLAO", "CALLAO"),
    ("ICA", "NASCA"): ("ICA", "NAZCA"),
}

FUENTE = "open_meteo"
GEOJSON = "https://raw.githubusercontent.com/juaneladio/peru-geojson/master/peru_provincial_simple.geojson"
API = "https://archive-api.open-meteo.com/v1/archive"
EXTRA = {}  # {("LORETO", "Putumayo"): ("1608", lat, lon)}

IGP = "igp"
IGP_DATASET_ID = "35ad4b2c-0084-4c68-a649-e2aff498bcc0"
IGP_API = "https://www.datosabiertos.gob.pe/api/3/action/package_show"

INDECI = "indeci"
INDECI_URL = "https://www.datosabiertos.gob.pe/sites/default/files/BD_2003-2025_EMERGENCIAS.csv"
INDECI_FUENTE = "INDECI - Plataforma Nacional de Datos Abiertos (SINPAD)"
UMBRAL_NULOS_PCT = 30.0

INGEMMET = "ingemmet"
INGEMMET_DATASET_ID = "243f9058-d9df-4c5f-b2b9-bbe2676a4d33"
INGEMMET_FUENTE = "INGEMMET - Inventario de Peligros Geológicos a Nivel Nacional"
UMBRAL_FALLECIDOS_ATIPICO = 50  # un registro con más fallecidos se marca para revisión (no se borra)

PESOS = {
    "fallecidos": 10.0, "desaparecidos": 8.0, "heridos": 5.0, "damnificados": 3.0,
    "afectados": 1.0, "viviendas_destruidas": 4.0, "viviendas_afectadas": 1.5,
    "centros_educativos_destruidos": 3.0, "centros_salud_destruidos": 4.0,
    "has_cultivo_destruido": 0.5,
}
KW_IMPACTO = (
    "fallec", "muert", "herid", "lesion", "desap", "damnif", "afect",
    "viv", "destru", "inhab", "perd", "hect", "cultiv",
    "pers", "poblac", "puente", "carretera", "camino",
    "agua", "desague", "canal", "vacuno", "caballar", "auquen",
    "caprino", "porcino", "camelido", "animales", "aula",
    "peso_de_la_ayuda", "costo_de_la_ayuda", "scant", "safec",
)
MESES_ES = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SETIEMBRE": 9, "SEPTIEMBRE": 9,
    "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12,
}
FORMATOS_FECHA = [
    "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
    "%d-%m-%Y", "%Y/%m/%d", "%Y%m%d", "%d/%m/%y", "%m/%d/%Y", "%d.%m.%Y",
]

FECHA = datetime.now(ZoneInfo(TZ)).strftime("%Y-%m-%d")

META_COLS = {
    "Fuente": "fuente", "Licencia": "licencia", "Tipo": "tipo",
    "Departamento": "departamento", "Provincia": "provincia",
    "Ubigeo provincia": "ubigeo_provincia", "Punto de consulta": "punto_consulta",
    "Latitud consulta": "latitud_consulta", "Longitud consulta": "longitud_consulta",
    "Latitud celda del modelo": "latitud_celda_modelo",
    "Longitud celda del modelo": "longitud_celda_modelo", "Altitud (m)": "altitud_m",
    "Periodo": "periodo", "Zona horaria": "zona_horaria",
    "Humedad (%)": "humedad_descripcion", "Generado (Lima)": "generado_lima",
}
COLS = {
    "Año": "anio", "Mes": "mes", "Día": "dia", "Temp. Máx (°C)": "temp_max_c",
    "Temp. Mín (°C)": "temp_min_c", "Humedad (%)": "humedad_pct",
    "Precipitación (mm)": "precipitacion_mm",
}
NUM = ["temp_max_c", "temp_min_c", "humedad_pct", "precipitacion_mm"]
NUM_META = ["latitud_consulta", "longitud_consulta", "latitud_celda_modelo",
            "longitud_celda_modelo", "altitud_m"]
COLUMNAS = ["fecha", "anio", "mes", "dia"] + NUM
GOLD_COLS = ["fecha", "departamento", "provincia",
             "temp_max_c", "temp_min_c", "temp_prom_c", "humedad_pct", "precipitacion_mm"]
FORMATOS_OM = {"fecha": "DD-MM-YYYY", "generado_lima": "YYYY-MM-DD HH:MM:SS"}
FORMATOS_IGP = {"FECHA_UTC": "DD-MM-YYYY", "HORA_UTC": "HH:MM:SS",
                "TIMESTAMP_EVENTO_UTC": "DD-MM-YYYY HH:MM:SS", "FECHA_CORTE": "DD-MM-YYYY",
                "FECHA_CARGA_SILVER": "DD-MM-YYYY HH:MM:SS"}

SCOPE_SECRETOS = "amankay"        # secret scope de Databricks con las 3 llaves de AWS
ESPERA_CREDENCIALES_MIN = 10      # si vencen, cuántos minutos espera a que actualices los secretos
ACCESS_KEY = SECRET_KEY = SESSION_TOKEN = None
s3 = None

try:  # Databricks: el bucket se elige con un widget (arriba del notebook)
    dbutils.widgets.text("bucket", BUCKET_DEFECTO, "Bucket S3")
    BUCKET = dbutils.widgets.get("bucket").strip() or BUCKET_DEFECTO
except NameError:  # fuera de Databricks (Colab o local)
    BUCKET = input(f"Nombre del bucket [{BUCKET_DEFECTO}]: ").strip() or BUCKET_DEFECTO


def slug(t):
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", t.lower()).strip("_")


def snake(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^0-9a-zA-Z]+", "_", s).strip("_").lower()


def norm_txt(x):
    if pd.isna(x):
        return x
    x = unicodedata.normalize("NFKD", str(x)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", x).strip().upper()


def conectar():
    args = dict(aws_access_key_id=ACCESS_KEY, aws_secret_access_key=SECRET_KEY,
                aws_session_token=SESSION_TOKEN, region_name=REGION)
    try:
        ident = boto3.client("sts", **args).get_caller_identity()
    except ClientError as e:
        raise RuntimeError(f"Credenciales inválidas o vencidas ({e.response['Error']['Code']}). "
                           "Revisa que las 3 llaves (access key, secret y token) sean de la misma sesión.") from None
    print("Credenciales OK:", ident["Arn"])
    cliente = boto3.client("s3", **args)
    try:
        cliente.list_objects_v2(Bucket=BUCKET, MaxKeys=1)
    except ClientError as e:
        raise RuntimeError(f"Credenciales válidas, pero sin acceso al bucket '{BUCKET}' "
                           f"({e.response['Error']['Code']}). Revisa el nombre y los permisos del rol.") from None
    return cliente


CODIGOS_VENCIDO = {"ExpiredToken", "ExpiredTokenException", "RequestExpired", "InvalidToken", "TokenRefreshRequired"}
_candado_aws = threading.Lock()


def leer_credenciales():
    """Databricks: lee los secretos del scope. Fuera de Databricks: las pide con getpass."""
    try:
        g = lambda k: dbutils.secrets.get(SCOPE_SECRETOS, k).strip()
    except NameError:
        return (getpass("Access key ID: ").strip(), getpass("Secret access key: ").strip(),
                getpass("Session token (Enter si no tienes): ").strip() or None)
    try:
        token = g("aws_session_token") or None
    except Exception:
        token = None
    return g("aws_access_key"), g("aws_secret_key"), token


def renovar_credenciales():
    """Lee las credenciales y reconecta. Si vencen a mitad de la corrida, espera a que actualices los secretos."""
    global ACCESS_KEY, SECRET_KEY, SESSION_TOKEN, s3
    with _candado_aws:
        probadas, avisado = None, False
        limite = time.time() + ESPERA_CREDENCIALES_MIN * 60
        while True:
            nuevas = leer_credenciales()
            if nuevas != probadas:
                probadas = nuevas
                ACCESS_KEY, SECRET_KEY, SESSION_TOKEN = nuevas
                try:
                    s3 = conectar()
                    print(f"Conectado al bucket '{BUCKET}'. Continuando...\n")
                    return
                except RuntimeError as e:
                    print(e)
            try:
                dbutils
            except NameError:
                probadas = None  # fuera de Databricks: vuelve a pedirlas
                continue
            if not avisado:
                print(f"\nActualiza los secretos del scope '{SCOPE_SECRETOS}' (aws_access_key, aws_secret_key, "
                      f"aws_session_token) con las credenciales nuevas.\n"
                      f"Esperando hasta {ESPERA_CREDENCIALES_MIN} min; el pipeline continúa solo en cuanto cambien...")
                avisado = True
            if time.time() > limite:
                raise RuntimeError("Credenciales AWS vencidas y no se actualizaron los secretos a tiempo. "
                                   "Actualízalos y vuelve a ejecutar la celda: continúa donde se quedó.")
            time.sleep(15)


def aws(fn):
    """Si las credenciales no existen o vencieron, las vuelve a pedir y repite la llamada (el avance no se pierde)."""
    @functools.wraps(fn)
    def envoltura(*a, **k):
        for _ in range(6):
            if s3 is None:
                renovar_credenciales()
            try:
                return fn(*a, **k)
            except ClientError as e:
                if e.response["Error"]["Code"] not in CODIGOS_VENCIDO:
                    raise
                print("\nLas credenciales de AWS vencieron.")
                renovar_credenciales()
        raise RuntimeError("Credenciales AWS vencidas repetidamente")
    return envoltura


@aws
def listar(prefijo):
    claves = []
    for pag in s3.get_paginator("list_objects_v2").paginate(Bucket=BUCKET, Prefix=prefijo):
        claves += [o["Key"] for o in pag.get("Contents", [])]
    return claves


@aws
def existe(clave):
    try:
        s3.head_object(Bucket=BUCKET, Key=clave)
        return True
    except ClientError as e:
        if e.response["Error"]["Code"] in ("404", "NoSuchKey", "NotFound"):
            return False
        raise


@aws
def leer_bytes(clave):
    return s3.get_object(Bucket=BUCKET, Key=clave)["Body"].read()


def leer(clave):
    return leer_bytes(clave).decode("utf-8")


@aws
def subir(clave, texto, tipo="text/csv; charset=utf-8"):
    s3.put_object(Bucket=BUCKET, Key=clave, Body=texto.encode("utf-8"), ContentType=tipo)


@aws
def subir_bytes(clave, datos, tipo):
    s3.put_object(Bucket=BUCKET, Key=clave, Body=datos, ContentType=tipo)


def subir_json(clave, obj):
    subir(clave, json.dumps(obj, ensure_ascii=False, indent=2, default=str), "application/json; charset=utf-8")


def ultimos(capa, ext):
    filtro = {slug(d) for d in DEPARTAMENTOS} if DEPARTAMENTOS else None
    mejores = {}
    for k in listar(f"{capa}/{FUENTE}/"):
        if not k.endswith(ext):
            continue
        _, _, dep, prov, fecha, _ = k.split("/")
        if filtro and dep not in filtro:
            continue
        if (dep, prov) not in mejores or fecha > mejores[(dep, prov)][0]:
            mejores[(dep, prov)] = (fecha, k)
    return dict(sorted(mejores.items()))


def ultimo_simple(capa, fuente, ext, contiene=""):
    mejor = None
    for k in listar(f"{capa}/{fuente}/"):
        if not k.endswith(ext) or contiene not in basename(k):
            continue
        fecha = k.split("/")[2]
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", fecha):
            continue  # ignora carpetas que no son una fecha (p. ej. la antigua .../reporte/noticia/)
        if mejor is None or fecha > mejor[0]:
            mejor = (fecha, k)
    if mejor is None:
        raise FileNotFoundError(f"No hay archivos {ext} en s3://{BUCKET}/{capa}/{fuente}/")
    return mejor


def num(x):
    return None if pd.isna(x) else float(x)


def completitud(df):
    return round(float((1 - df.isna().mean().mean()) * 100), 2)


def a_xlsx(df, formatos):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="datos")
        ws = w.sheets["datos"]
        ws.freeze_panes = "A2"
        for j, c in enumerate(df.columns, 1):
            s = df[c]
            derecha = (c in formatos
                       or (pd.api.types.is_numeric_dtype(s) and not pd.api.types.is_bool_dtype(s))
                       or pd.api.types.is_datetime64_any_dtype(s))
            alin = Alignment(horizontal="right" if derecha else "left")
            letra = get_column_letter(j)
            for celda in ws[letra]:
                celda.alignment = alin
            if c in formatos:
                for celda in ws[letra][1:]:
                    celda.number_format = formatos[c]
            largo = max([len(str(c))] + s.astype(str).str.len().tolist()[:200])
            ws.column_dimensions[letra].width = min(largo + 2, 60)
    return buf.getvalue()


# ======================= OPEN-METEO =======================

def parsear(texto):
    lineas = texto.splitlines()
    meta = [l for l in lineas if l.startswith("# ")]
    df = pd.read_csv(io.StringIO("\n".join(lineas[len(meta):])), sep=SEP,
                     na_values=["S/D"], dtype={"Año": str, "Mes": str, "Día": str})
    return meta, df


def cargar_provincias():
    d = requests.get(GEOJSON, timeout=60).json()
    grupos = defaultdict(list)
    for f in d["features"]:
        grupos[f["properties"]["FIRST_IDPR"]].append(f)
    provincias = []
    for codigo, fs in grupos.items():
        fs.sort(key=lambda f: f["properties"]["ha"], reverse=True)
        punto = unary_union([shape(f["geometry"]) for f in fs]).representative_point()
        provincias.append({
            "dep": fs[0]["properties"]["FIRST_NOMB"], "prov": fs[0]["properties"]["NOMBPROV"],
            "codigo": codigo, "lat": round(punto.y, 5), "lon": round(punto.x, 5),
        })
    for (dep, prov), (codigo, lat, lon) in EXTRA.items():
        provincias.append({"dep": dep.upper(), "prov": prov.upper(), "codigo": codigo, "lat": lat, "lon": lon})
    return sorted(provincias, key=lambda p: (p["dep"], p["prov"]))


def cargar_poligonos():
    d = requests.get(GEOJSON, timeout=60).json()
    geoms = [shape(f["geometry"]) for f in d["features"]]
    nombres = [(f["properties"]["FIRST_NOMB"], f["properties"]["NOMBPROV"]) for f in d["features"]]
    return geoms, nombres


def pedir(lat, lon, ini, fin):
    params = {
        "latitude": lat, "longitude": lon,
        "start_date": ini.isoformat(), "end_date": fin.isoformat(),
        "daily": "temperature_2m_max,temperature_2m_min,precipitation_sum",
        "hourly": "relative_humidity_2m",
        "timezone": TZ,
    }
    for intento in range(5):
        r = requests.get(API, params=params, timeout=60)
        if r.status_code == 200:
            return r.json()
        if r.status_code == 429 and "daily" in r.text.lower():
            raise RuntimeError("Límite diario de Open-Meteo alcanzado")
        if r.status_code in (429, 500, 502, 503, 504):
            time.sleep(10 * (intento + 1))
            continue
        raise RuntimeError(f"{r.status_code}: {r.text[:200]}")
    raise RuntimeError("Reintentos agotados")


def tramos(ini, fin):
    a = ini
    while a <= fin:
        b = min(date(a.year, 12, 31), fin)
        yield a, b
        a = b + timedelta(days=1)


def a_dataframe(js):
    d = pd.DataFrame(js["daily"])
    h = pd.DataFrame(js["hourly"])
    h["fecha"] = h["time"].str[:10]
    d["humedad"] = d["time"].map(h.groupby("fecha")["relative_humidity_2m"].mean().round(1))
    return d


def formato_senamhi(df):
    f = pd.to_datetime(df["time"])
    return pd.DataFrame({
        "Año": f.dt.strftime("%Y"),
        "Mes": f.dt.strftime("%m"),
        "Día": f.dt.strftime("%d"),
        "Temp. Máx (°C)": df["temperature_2m_max"].values,
        "Temp. Mín (°C)": df["temperature_2m_min"].values,
        "Humedad (%)": df["humedad"].values,
        "Precipitación (mm)": df["precipitation_sum"].values,
    })


def metadata(p, js, ini, fin):
    ahora = datetime.now(ZoneInfo(TZ)).strftime("%Y-%m-%d %H:%M:%S")
    return [
        "# Fuente: Open-Meteo Historical Weather API (reanálisis ERA5/ERA5-Land)",
        "# Licencia: CC BY 4.0 - Open-Meteo.com",
        "# Tipo: Reanálisis (dato modelado, no es una estación)",
        f"# Departamento: {p['dep']}",
        f"# Provincia: {p['prov'].title()}",
        f"# Ubigeo provincia: {p['codigo']}",
        "# Punto de consulta: punto representativo dentro del límite de la provincia",
        f"# Latitud consulta: {p['lat']}",
        f"# Longitud consulta: {p['lon']}",
        f"# Latitud celda del modelo: {js['latitude']}",
        f"# Longitud celda del modelo: {js['longitude']}",
        f"# Altitud (m): {js.get('elevation')}",
        f"# Periodo: {ini} a {fin}",
        f"# Zona horaria: {TZ}",
        "# Humedad (%): promedio diario de las 24 horas",
        f"# Generado (Lima): {ahora}",
    ]


def descargar_provincia(p, fin_real, clave):
    partes, js = [], None
    for a, b in tramos(INICIO, fin_real):
        js = pedir(p["lat"], p["lon"], a, b)
        partes.append(a_dataframe(js))
        time.sleep(PAUSA)
    salida = formato_senamhi(pd.concat(partes, ignore_index=True))
    texto = "\n".join(metadata(p, js, INICIO, fin_real))
    texto += "\n" + salida.to_csv(sep=SEP, index=False, na_rep="S/D")
    return texto, len(salida)  # se sube desde el hilo principal (ver om_fuente)


def om_fuente():
    provincias = [p for p in cargar_provincias() if not DEPARTAMENTOS or p["dep"] in DEPARTAMENTOS]
    fin_real = min(FIN, datetime.now(ZoneInfo(TZ)).date() - timedelta(days=7))
    etiqueta = str(INICIO.year) if INICIO.year == fin_real.year else f"{INICIO.year}_{fin_real.year}"
    existentes = set(listar(f"{CAPA_FUENTE}/{FUENTE}/"))

    pendientes = []
    for p in provincias:
        clave = f"{CAPA_FUENTE}/{FUENTE}/{slug(p['dep'])}/{slug(p['prov'])}/{FECHA}/{slug(p['prov'])}_{etiqueta}.csv"
        if clave not in existentes:
            pendientes.append((p, clave))
    print(f"{len(pendientes)} provincias por descargar ({len(provincias) - len(pendientes)} ya estaban).")

    ok, fallidos = 0, []
    with ThreadPoolExecutor(HILOS) as ex:
        futuros = {ex.submit(descargar_provincia, p, fin_real, clave): (p, clave)
                   for p, clave in pendientes}
        for i, f in enumerate(as_completed(futuros), 1):
            p, clave = futuros[f]
            nombre = f"{p['dep']} / {p['prov'].title()}"
            try:
                texto, filas = f.result()
                subir(clave, texto)
                ok += 1
                print(f"[{i}/{len(pendientes)}] {nombre}: {filas} filas -> s3://{BUCKET}/{clave}")
            except Exception as e:
                fallidos.append(nombre)
                print(f"[{i}/{len(pendientes)}] {nombre} falló: {e!r}")
    print(f"Fuente: {ok} archivos subidos, {len(fallidos)} fallidos {fallidos or ''}")


def om_bronze():
    por_dep = defaultdict(list)
    for (dep, prov), (f_fuente, clave) in ultimos(CAPA_FUENTE, ".csv").items():
        texto = leer(clave)
        meta, df = parsear(texto)
        nombre = f"{basename(clave)[:-4]}_{FECHA}.csv"
        destino = f"bronze/{FUENTE}/{dep}/{prov}/{FECHA}/{nombre}"
        subir(destino, texto)
        por_dep[dep].append({
            "provincia": prov,
            "archivo_origen": clave,
            "fecha_origen": f_fuente,
            "archivo_destino": destino,
            "filas_metadata": len(meta),
            "filas_datos": len(df),
            "columnas": list(df.columns),
            "columna_final": df.columns[-1],
            "primera_fecha": f"{df.iloc[0]['Año']}-{df.iloc[0]['Mes']}-{df.iloc[0]['Día']}",
            "ultima_fecha": f"{df.iloc[-1]['Año']}-{df.iloc[-1]['Mes']}-{df.iloc[-1]['Día']}",
            "ultimo_dato": {c: (None if pd.isna(v) else str(v)) for c, v in df.iloc[-1].items()},
            "bytes": len(texto.encode("utf-8")),
            "sha256": hashlib.sha256(texto.encode("utf-8")).hexdigest(),
        })
        print(f"bronze: {dep}/{prov}")
    for dep, items in por_dep.items():
        subir_json(f"bronze/{FUENTE}/{dep}/{dep}_trazabilidad_{FECHA}.json", {
            "capa": "bronze", "fuente": FUENTE, "departamento": dep,
            "fecha_proceso": FECHA, "zona_horaria": TZ,
            "provincias": len(items), "detalle": items,
        })


REGLAS = {
    "R01_fecha_invalida": "Se eliminan filas cuya fecha (año-mes-día) no es válida",
    "R02_fecha_duplicada": "Se elimina la fila repetida de una misma fecha (se conserva la primera)",
    "R03_temp_max_rango": "temp_max_c fuera de [-30, 50] se anula",
    "R04_temp_min_rango": "temp_min_c fuera de [-30, 50] se anula",
    "R05_temp_min_mayor_max": "temp_min_c mayor que temp_max_c se anula",
    "R06_humedad_rango": "humedad_pct fuera de [0, 100] se anula",
    "R07_precipitacion_rango": "precipitacion_mm fuera de [0, 500] se anula",
}


def calidad(df):
    n = len(df)
    return {
        "filas": n,
        "fechas_duplicadas": int(df.duplicated(["anio", "mes", "dia"]).sum()),
        "nulos": {c: int(df[c].isna().sum()) for c in NUM},
        "pct_nulos": {c: round(float(df[c].isna().mean() * 100), 2) if n else None for c in NUM},
        "minimo": {c: num(df[c].min()) for c in NUM},
        "maximo": {c: num(df[c].max()) for c in NUM},
        "promedio": {c: num(round(df[c].mean(), 2)) for c in NUM},
    }


def aplicar_reglas(df):
    df = df.copy()
    afectadas = {k: 0 for k in REGLAS}
    for c in NUM:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df["fecha"] = pd.to_datetime(df["anio"] + "-" + df["mes"] + "-" + df["dia"],
                                 format="%Y-%m-%d", errors="coerce")

    malas = df["fecha"].isna()
    afectadas["R01_fecha_invalida"] = int(malas.sum())
    df = df[~malas]

    dup = df["fecha"].duplicated(keep="first")
    afectadas["R02_fecha_duplicada"] = int(dup.sum())
    df = df[~dup].copy()

    for regla, col, lo, hi in [("R03_temp_max_rango", "temp_max_c", -30, 50),
                               ("R04_temp_min_rango", "temp_min_c", -30, 50),
                               ("R06_humedad_rango", "humedad_pct", 0, 100),
                               ("R07_precipitacion_rango", "precipitacion_mm", 0, 500)]:
        m = (df[col] < lo) | (df[col] > hi)
        afectadas[regla] = int(m.sum())
        df.loc[m, col] = pd.NA
        df[col] = pd.to_numeric(df[col])

    m = df["temp_min_c"] > df["temp_max_c"]
    afectadas["R05_temp_min_mayor_max"] = int(m.sum())
    df.loc[m, "temp_min_c"] = pd.NA
    df["temp_min_c"] = pd.to_numeric(df["temp_min_c"])

    for c in ["anio", "mes", "dia"]:
        df[c] = df[c].astype(int)
    return df.reset_index(drop=True), afectadas


def om_silver():
    por_dep = defaultdict(list)
    for (dep, prov), (f_bronze, clave) in ultimos("bronze", ".csv").items():
        meta, df = parsear(leer(clave))
        df = df.rename(columns=COLS)
        m = {}
        for l in meta:
            k, v = l[2:].split(": ", 1)
            m[META_COLS.get(k, slug(k))] = v

        antes = calidad(df)
        df, afectadas = aplicar_reglas(df)
        despues = calidad(df)

        for c, v in m.items():
            df[c] = v
        for c in NUM_META:
            df[c] = pd.to_numeric(df[c], errors="coerce")
        df["generado_lima"] = pd.to_datetime(df["generado_lima"], errors="coerce")
        df = df[list(m) + COLUMNAS]

        stem = basename(clave)[:-4]
        carpeta = f"silver/{FUENTE}/{dep}/{prov}/{FECHA}"
        subir_bytes(f"{carpeta}/{stem}.xlsx", a_xlsx(df, FORMATOS_OM), XLSX)
        reglas = [{"regla": k, "descripcion": REGLAS[k], "filas_afectadas": afectadas[k]} for k in REGLAS]
        subir_json(f"{carpeta}/{stem}_calidad.json", {
            "provincia": prov, "archivo_origen": clave, "fecha_proceso": FECHA,
            "antes": antes, "despues": despues, "reglas_aplicadas": reglas,
        })
        por_dep[dep].append({
            "provincia": prov,
            "archivo_origen": clave,
            "archivo_destino": f"{carpeta}/{stem}.xlsx",
            "archivo_calidad": f"{carpeta}/{stem}_calidad.json",
            "filas_entrada": antes["filas"],
            "filas_salida": despues["filas"],
            "columnas_renombradas": {**{f"# {k}": v for k, v in META_COLS.items()}, **COLS},
            "columnas_silver": list(df.columns),
            "columna_final": df.columns[-1],
            "ultima_fecha": df["fecha"].iloc[-1].strftime("%Y-%m-%d") if len(df) else None,
            "reglas_afectadas": {r["regla"]: r["filas_afectadas"] for r in reglas},
        })
        print(f"silver: {dep}/{prov}")
    for dep, items in por_dep.items():
        subir_json(f"silver/{FUENTE}/{dep}/{dep}_trazabilidad_{FECHA}.json", {
            "capa": "silver", "fuente": FUENTE, "departamento": dep,
            "fecha_proceso": FECHA, "zona_horaria": TZ,
            "provincias": len(items), "detalle": items,
        })


def om_gold():
    partes = []
    for (dep, prov), (f_silver, clave) in ultimos("silver", ".xlsx").items():
        partes.append(pd.read_excel(io.BytesIO(leer_bytes(clave)), sheet_name="datos"))
    df = pd.concat(partes, ignore_index=True)
    df["temp_prom_c"] = ((df["temp_max_c"] + df["temp_min_c"]) / 2).round(1)
    df = df.sort_values(["departamento", "provincia", "fecha"])
    anios = df["fecha"].dt.year
    etiqueta = str(anios.min()) if anios.min() == anios.max() else f"{anios.min()}_{anios.max()}"
    df["fecha"] = df["fecha"].dt.strftime("%d-%m-%Y")
    df = df[GOLD_COLS]
    clave = f"gold/{FUENTE}/{FECHA}/{FUENTE}_{etiqueta}_{FECHA}.csv"
    subir(clave, df.to_csv(sep=SEP, index=False, na_rep=""))
    print(f"gold: {len(df)} filas -> s3://{BUCKET}/{clave}")


# ============================ IGP ============================

def igp_fuente():
    clave = f"{CAPA_FUENTE}/{IGP}/{FECHA}/igp_sismos.csv"
    if existe(clave):
        print(f"Ya existe en S3, no se descarga: {clave}")
        return
    r = requests.get(IGP_API, params={"id": IGP_DATASET_ID}, headers=UA, timeout=30)
    r.raise_for_status()
    res = r.json()["result"]
    dataset = res[0] if isinstance(res, list) else res
    csvs = [x for x in dataset["resources"]
            if x.get("format", "").lower().lstrip(".") == "csv" and x.get("url")]
    if not csvs:
        raise RuntimeError("El dataset del IGP no tiene recursos CSV con URL")
    recurso = csvs[0]
    print("Descargando:", recurso["url"])
    archivo = requests.get(recurso["url"], headers=UA, timeout=120)
    archivo.raise_for_status()

    ahora = datetime.now(ZoneInfo(TZ))
    filas = list(csv.reader(io.StringIO(archivo.content.decode("utf-8", errors="replace"))))
    subir_bytes(clave, archivo.content, "text/csv; charset=utf-8")
    subir_json(clave[:-4] + ".json", {
        "fuente": IGP,
        "dataset": dataset.get("title"),
        "url_origen": recurso["url"],
        "fecha_modificacion_fuente": dataset.get("metadata_modified"),
        "fecha_hora_extraccion": ahora.isoformat(),
        "cantidad_registros": len(filas) - 1,
        "formato": "csv",
        "ubicacion": f"s3://{BUCKET}/{clave}",
        "estado": "extraido",
    })
    print(f"fuente igp: {len(filas) - 1} registros -> s3://{BUCKET}/{clave}")


def igp_bronze():
    f_fuente, clave = ultimo_simple(CAPA_FUENTE, IGP, ".csv")
    datos = leer_bytes(clave)
    ficha_fuente = json.loads(leer(clave[:-4] + ".json"))
    df = pd.read_csv(io.BytesIO(datos), sep=SEP, low_memory=False)
    nombre = f"igp_sismos_{FECHA}"
    destino = f"bronze/{IGP}/{FECHA}/{nombre}.csv"
    subir_bytes(destino, datos, "text/csv; charset=utf-8")
    subir_json(f"bronze/{IGP}/{FECHA}/{nombre}.json", {
        "capa": "bronze", "fuente": IGP, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "archivo_origen": clave, "fecha_origen": f_fuente,
        "fecha_hora_extraccion": ficha_fuente.get("fecha_hora_extraccion"),
        "fecha_modificacion_fuente": ficha_fuente.get("fecha_modificacion_fuente"),
        "archivo_destino": destino,
        "filas_datos": len(df),
        "columnas": list(df.columns),
        "columna_final": df.columns[-1],
        "bytes": len(datos),
        "sha256": hashlib.sha256(datos).hexdigest(),
        "estado": "copiado",
    })
    print(f"bronze igp: {len(df)} filas -> s3://{BUCKET}/{destino}")


def calidad_igp(df):
    r = {
        "filas": len(df),
        "completitud_pct": completitud(df),
        "nulos_por_columna": {c: int(n) for c, n in df.isna().sum().items() if n},
    }
    if "DENTRO_RANGO_PERU" in df:
        r["latitud_en_rango_pct"] = round(float(df["DENTRO_RANGO_PERU"].mean() * 100), 2)
        r["profundidad_valida_pct"] = round(float(df["PROFUNDIDAD_VALIDA"].mean() * 100), 2)
    return r


def igp_silver():
    f_bronze, clave = ultimo_simple("bronze", IGP, ".csv")
    df = pd.read_csv(io.BytesIO(leer_bytes(clave)), sep=SEP, low_memory=False)
    n_original = len(df)
    cols_bronze = list(df.columns)
    tipos_bronze = df.dtypes.astype(str).to_dict()
    antes = calidad_igp(df)
    reglas = []

    def reg(accion, campo, detalle, motivo, fa, fd):
        reglas.append({"accion": accion, "campo": campo, "detalle": detalle, "motivo": motivo,
                       "filas_antes": fa, "filas_despues": fd})

    hora_str = df["HORA_UTC"].astype(int).astype(str).str.zfill(6)
    fecha_str = df["FECHA_UTC"].astype(int).astype(str).str.zfill(8)
    df["FECHA_UTC"] = pd.to_datetime(fecha_str, format="%Y%m%d")
    reg("convertir_tipo", "FECHA_UTC", "entero YYYYMMDD a fecha", "permitir operar por fecha", n_original, len(df))
    df["HORA_UTC"] = pd.to_datetime(hora_str, format="%H%M%S").dt.time
    reg("convertir_tipo", "HORA_UTC", "entero sin ceros a HH:MM:SS", "se perdían los ceros a la izquierda", n_original, len(df))
    df["TIMESTAMP_EVENTO_UTC"] = pd.to_datetime(fecha_str + hora_str, format="%Y%m%d%H%M%S")
    reg("agregar_columna", "TIMESTAMP_EVENTO_UTC", "fecha + hora combinadas", "facilitar ordenar por evento", n_original, len(df))

    a = len(df)
    df = df[df["FECHA_UTC"].between(pd.Timestamp(INICIO), pd.Timestamp(FIN))].copy()
    reg("filtrar_filas", "FECHA_UTC", f"{a - len(df)} filas fuera de {INICIO} a {FIN}", "solo el periodo solicitado", a, len(df))
    if df.empty:
        raise ValueError(f"El catálogo del IGP no tiene sismos entre {INICIO} y {FIN}")

    df["DENTRO_RANGO_PERU"] = df["LATITUD"].between(-20, 2)
    reg("agregar_flag", "LATITUD", f"{int((~df['DENTRO_RANGO_PERU']).sum())} filas fuera del rango -20 a 2", "marcar en vez de eliminar", len(df), len(df))
    df["PROFUNDIDAD_VALIDA"] = df["PROFUNDIDAD"].between(0, 700)
    reg("agregar_flag", "PROFUNDIDAD", f"{int((~df['PROFUNDIDAD_VALIDA']).sum())} filas fuera de 0-700 km", "marcar en vez de eliminar", len(df), len(df))

    invertidas = [0]

    def normaliza_fecha_corte(x):
        s = str(int(x)).zfill(8)
        y, p, q = s[0:4], s[4:6], s[6:8]
        try:
            return pd.Timestamp(year=int(y), month=int(p), day=int(q))
        except ValueError:
            invertidas[0] += 1
            return pd.Timestamp(year=int(y), month=int(q), day=int(p))

    df["FECHA_CORTE"] = df["FECHA_CORTE"].apply(normaliza_fecha_corte)
    reg("normalizar", "FECHA_CORTE", f"{invertidas[0]} fechas con día/mes invertido corregidas", "formato mixto", len(df), len(df))

    comparar = ["FECHA_UTC", "HORA_UTC", "LATITUD", "LONGITUD", "PROFUNDIDAD", "MAGNITUD"]
    a = len(df)
    df = df.sort_values("ID").drop_duplicates(subset=comparar, keep="first")
    n_final = len(df)
    reg("eliminar_filas", "fila completa", f"{a - n_final} duplicados eliminados", "mismo evento con dos ID, se conserva el menor", a, n_final)

    df["ARCHIVO_ORIGEN"] = clave
    df["FECHA_CARGA_SILVER"] = pd.Timestamp(datetime.now(ZoneInfo(TZ)).replace(tzinfo=None))

    errores = []
    if df["ID"].duplicated().any():
        errores.append("hay ID duplicados")
    for col in ["ID", "FECHA_UTC", "MAGNITUD"]:
        nulos = int(df[col].isna().sum())
        if nulos:
            errores.append(f"{col} tiene {nulos} nulos")
    if errores:
        raise ValueError("Validación fallida: " + "; ".join(errores))
    reg("validar", "ID, FECHA_UTC, MAGNITUD", "sin duplicados en ID y sin nulos", "control de calidad antes de subir", n_final, n_final)

    cols_silver = list(df.columns)
    eliminadas = [c for c in cols_bronze if c not in cols_silver]
    agregadas = [c for c in cols_silver if c not in cols_bronze]
    despues = calidad_igp(df)

    stem = basename(clave)[:-4]
    carpeta = f"silver/{IGP}/{FECHA}"
    subir_bytes(f"{carpeta}/{stem}.xlsx", a_xlsx(df, FORMATOS_IGP), XLSX)
    subir_json(f"{carpeta}/{stem}_calidad.json", {
        "archivo_origen": clave, "fecha_proceso": FECHA,
        "antes": antes, "despues": despues, "reglas_aplicadas": reglas,
    })
    subir_json(f"{carpeta}/{stem}.json", {
        "capa": "silver", "fuente": IGP, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "archivo_origen": clave, "fecha_origen": f_bronze,
        "archivo_destino": f"{carpeta}/{stem}.xlsx",
        "archivo_calidad": f"{carpeta}/{stem}_calidad.json",
        "periodo": f"{INICIO} a {FIN}",
        "filas_entrada": n_original, "filas_salida": n_final,
        "columnas_eliminadas": eliminadas, "columnas_agregadas": agregadas,
        "columna_final": cols_silver[-1],
        "esquema": [{"columna": c, "tipo_bronze": tipos_bronze.get(c, "-"), "tipo_silver": str(df[c].dtype)}
                    for c in cols_silver],
        "estado": "procesado",
    })
    print(f"silver igp: {n_original} -> {n_final} filas -> s3://{BUCKET}/{carpeta}/{stem}.xlsx")


def igp_gold():
    f_silver, clave = ultimo_simple("silver", IGP, ".xlsx")
    df = pd.read_excel(io.BytesIO(leer_bytes(clave)), sheet_name="datos")
    df = df.drop(columns=["ARCHIVO_ORIGEN", "FECHA_CARGA_SILVER"])

    geoms, nombres = cargar_poligonos()
    arbol = STRtree(geoms)
    ok = df["LATITUD"].notna() & df["LONGITUD"].notna()
    idx = arbol.nearest(shapely.points(df.loc[ok, "LONGITUD"].to_numpy(), df.loc[ok, "LATITUD"].to_numpy()))
    df["DEPARTAMENTO_CERCANO"] = pd.NA
    df["PROVINCIA_CERCANA"] = pd.NA
    df.loc[ok, "DEPARTAMENTO_CERCANO"] = [nombres[i][0] for i in idx]
    df.loc[ok, "PROVINCIA_CERCANA"] = [nombres[i][1] for i in idx]

    anios = df["FECHA_UTC"].dt.year
    etiqueta = str(anios.min()) if anios.min() == anios.max() else f"{anios.min()}_{anios.max()}"
    df["FECHA_UTC"] = df["FECHA_UTC"].dt.strftime("%d-%m-%Y")
    df["FECHA_CORTE"] = df["FECHA_CORTE"].dt.strftime("%d-%m-%Y")
    df["TIMESTAMP_EVENTO_UTC"] = df["TIMESTAMP_EVENTO_UTC"].dt.strftime("%d-%m-%Y %H:%M:%S")
    destino = f"gold/{IGP}/{FECHA}/igp_{etiqueta}_{FECHA}.csv"
    subir(destino, df.to_csv(sep=SEP, index=False, na_rep=""))
    print(f"gold igp: {len(df)} filas -> s3://{BUCKET}/{destino}")


# =========================== INDECI ===========================

def buscar(d, *claves):
    cols = list(d.columns)
    for k in claves:
        if k in cols:
            return k
    for k in claves:
        for c in cols:
            if k in c.split("_") or c == k:
                return c
    matches = [(len(c), c) for c in cols for k in claves if k in c]
    return sorted(matches)[0][1] if matches else None


def parse_fecha(serie):
    if pd.api.types.is_datetime64_any_dtype(serie):
        return pd.to_datetime(serie)
    txt = serie.astype("string").str.strip().str.replace(r"\.0+$", "", regex=True)
    txt = txt.mask(txt.isin(["", "nan", "NaN", "None", "NaT", "<NA>", "0"]))
    out = pd.Series(pd.NaT, index=serie.index, dtype="datetime64[ns]")
    for variante in (txt, txt.str.slice(0, 10), txt.str.slice(0, 19)):
        for fmt in FORMATOS_FECHA:
            if out.notna().all():
                break
            out = out.fillna(pd.to_datetime(variante, format=fmt, errors="coerce"))
    n = pd.to_numeric(txt, errors="coerce")
    pend = out.isna() & n.notna()
    if pend.any():
        ms = pend & (n > 1e11)
        sg = pend & (n > 1e8) & (n <= 1e11)
        xl = pend & (n > 20000) & (n < 80000)
        out = out.fillna(pd.to_datetime(n.where(ms), unit="ms", errors="coerce"))
        out = out.fillna(pd.to_datetime(n.where(sg), unit="s", errors="coerce"))
        out = out.fillna(pd.to_datetime(n.where(xl), unit="D", origin="1899-12-30", errors="coerce"))
    return out


def indeci_fuente():
    clave = f"{CAPA_FUENTE}/{INDECI}/{FECHA}/indeci_emergencias.csv"
    if existe(clave):
        print(f"Ya existe en S3, no se descarga: {clave}")
        return
    print("Descargando:", INDECI_URL)
    r = requests.get(INDECI_URL, headers=UA_CHROME, timeout=300)
    r.raise_for_status()
    subir_bytes(clave, r.content, "text/csv")
    subir_json(clave[:-4] + ".json", {
        "fuente": INDECI_FUENTE,
        "dataset": "emergencias_danos_2003_2025",
        "url_origen": INDECI_URL,
        "periodo_cubierto": "2003-2025",
        "fecha_hora_extraccion": datetime.now(ZoneInfo(TZ)).isoformat(),
        "formato": "csv", "origen_sep": ";", "origen_encoding": "latin-1",
        "bytes": len(r.content),
        "sha256": hashlib.sha256(r.content).hexdigest(),
        "ubicacion": f"s3://{BUCKET}/{clave}",
        "estado": "extraido",
    })
    print(f"fuente indeci: {len(r.content):,} bytes -> s3://{BUCKET}/{clave}")


def indeci_bronze():
    f_fuente, clave = ultimo_simple(CAPA_FUENTE, INDECI, ".csv")
    origen = leer_bytes(clave)
    ficha_fuente = json.loads(leer(clave[:-4] + ".json"))
    df = pd.read_csv(io.BytesIO(origen), sep=";", encoding="latin-1", low_memory=False)
    datos = df.to_csv(index=False, sep=SEP).encode("utf-8")

    anios = None
    for c in df.columns:
        if str(c).strip().upper() in ("AÑO", "ANO", "ANIO"):
            v = pd.to_numeric(df[c], errors="coerce").dropna()
            if len(v):
                anios = {"min": int(v.min()), "max": int(v.max())}
            break

    nombre = f"indeci_emergencias_{FECHA}"
    destino = f"bronze/{INDECI}/{FECHA}/{nombre}.csv"
    subir_bytes(destino, datos, "text/csv; charset=utf-8")
    subir_json(f"bronze/{INDECI}/{FECHA}/{nombre}.json", {
        "capa": "bronze", "fuente": INDECI_FUENTE, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "archivo_origen": clave, "fecha_origen": f_fuente,
        "fecha_hora_extraccion": ficha_fuente.get("fecha_hora_extraccion"),
        "url_origen": ficha_fuente.get("url_origen"),
        "archivo_destino": destino,
        "origen_sep": ";", "origen_encoding": "latin-1", "encoding": "utf-8", "sep": SEP,
        "anios_detectados": anios,
        "filas_datos": int(len(df)), "n_columnas": int(len(df.columns)),
        "columnas": list(df.columns),
        "columna_final": df.columns[-1],
        "tipos_de_datos": {c: str(t) for c, t in df.dtypes.items()},
        "nulos_por_columna": {c: int(df[c].isna().sum()) for c in df.columns},
        "bytes_origen": len(origen), "bytes_bronze": len(datos),
        "sha256": hashlib.sha256(datos).hexdigest(),
        "estado": "copiado",
    })
    print(f"bronze indeci: {len(df):,} filas -> s3://{BUCKET}/{destino}")


def indeci_silver():
    f_bronze, clave = ultimo_simple("bronze", INDECI, ".csv")
    df = pd.read_csv(io.BytesIO(leer_bytes(clave)), sep=SEP, low_memory=False)
    n0 = len(df)
    cols_bronze = list(df.columns)
    antes = {"filas": n0, "completitud_media_pct": completitud(df)}
    reglas = []

    def reg(accion, campo, detalle, fa, fd):
        reglas.append({"accion": accion, "campo": campo, "detalle": detalle,
                       "filas_antes": fa, "filas_despues": fd})

    nuevos, vistos = [], {}
    for c in df.columns:
        n = snake(c)
        vistos[n] = vistos.get(n, 0) + 1
        nuevos.append(n if vistos[n] == 1 else f"{n}_{vistos[n]}")
    df.columns = nuevos
    df["archivo_origen"] = clave
    reg("normalizar_nombres", "columnas", f"{len(nuevos)} columnas a minúsculas sin caracteres especiales", n0, n0)

    roles = {
        "codigo": buscar(df, "codigo_de_emergencia_sinpad", "codigo_de_emergencia", "sinpad"),
        "anio": buscar(df, "ano", "anio"),
        "mes": buscar(df, "mes"),
        "cod_distrito": buscar(df, "cod_distrito", "codigo_distrito"),
        "departamento": buscar(df, "dpto", "departamento"),
        "provincia": buscar(df, "prov", "provincia"),
        "distrito": buscar(df, "distrito", "dist"),
        "fenomeno": buscar(df, "peligro", "fenomeno", "evento"),
        "tipo_peligro": buscar(df, "tipo_de_peligro", "tipo_peligro"),
        "region_natural": buscar(df, "region_natural"),
    }
    cands = [c for c in df.columns if "fecha" in c or "date" in c]
    if not cands:
        raise ValueError("No se encontró columna de fecha en INDECI")
    parseadas = {c: int(parse_fecha(df[c]).notna().sum()) for c in cands}
    orden = sorted(cands, key=lambda c: -parseadas[c])
    F = orden[0]
    roles["fecha"] = F
    excl = {v for v in roles.values() if v} | {"archivo_origen"}
    impacto = [c for c in df.columns if c not in excl and any(k in c for k in KW_IMPACTO)]
    texto_cols = [c for c in df.columns
                  if (pd.api.types.is_object_dtype(df[c]) or pd.api.types.is_string_dtype(df[c]))
                  and c not in orden and c != "archivo_origen"]

    s = df.copy()
    for c in texto_cols:
        s[c] = s[c].map(norm_txt)
    reg("normalizar_texto", f"{len(texto_cols)} columnas", "MAYÚSCULAS, sin tildes ni espacios repetidos", n0, n0)

    fecha = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    for c in orden:
        fecha = fecha.fillna(parse_fecha(s[c]))
    s[F] = fecha.dt.normalize()
    hoy = pd.Timestamp(datetime.now(ZoneInfo(TZ)).replace(tzinfo=None))
    fuera = s[F].notna() & ((s[F] < "1990-01-01") | (s[F] > hoy))
    s.loc[fuera, F] = pd.NaT
    a = len(s)
    s = s[s[F].notna()]
    reg("eliminar_filas", F, f"{int(fuera.sum())} fuera de rango 1990-hoy; {a - len(s)} sin fecha válida eliminadas (incluye las fuera de rango)", a, len(s))

    a = len(s)
    s = s[(s[F] >= pd.Timestamp(INICIO)) & (s[F] < pd.Timestamp(FIN) + pd.Timedelta(days=1))].reset_index(drop=True)
    reg("filtrar_filas", F, f"{a - len(s)} filas fuera de {INICIO} a {FIN}", a, len(s))
    if s.empty:
        raise ValueError(f"INDECI no tiene emergencias entre {INICIO} y {FIN}")

    if impacto:
        for c in impacto:
            s[c] = pd.to_numeric(s[c], errors="coerce")
        nulas = int(s[impacto].isna().sum().sum())
        neg = int((s[impacto] < 0).sum().sum())
        s[impacto] = s[impacto].clip(lower=0).fillna(0)
        reg("imputar_valores", f"{len(impacto)} columnas de impacto", f"{nulas} celdas nulas y {neg} negativas a 0", len(s), len(s))

    for rol in ("departamento", "provincia", "distrito", "fenomeno", "tipo_peligro"):
        c = roles.get(rol)
        if c and c in s.columns:
            s[c] = s[c].astype("string")
            vac = s[c].isna() | (s[c].str.strip() == "") | s[c].str.upper().isin(["NAN", "<NA>", "NONE", "NAT"])
            n_vac = int(vac.sum())
            s.loc[vac, c] = "NO ESPECIFICADO"
            reg("imputar_valores", c, f"{n_vac} nulos o vacíos a 'NO ESPECIFICADO' ({rol})", len(s), len(s))

    cc = roles["cod_distrito"]
    if cc and cc in s.columns:
        s[cc] = s[cc].astype("string").str.strip().str.replace(r"\.0$", "", regex=True).str.zfill(6)
        reg("convertir_tipo", cc, "ubigeo como texto de 6 dígitos con ceros a la izquierda", len(s), len(s))

    cols_dup = [c for c in s.columns if c != "archivo_origen"]
    a = len(s)
    s = s.drop_duplicates(subset=cols_dup).reset_index(drop=True)
    reg("eliminar_filas", "fila completa", f"{a - len(s)} duplicados eliminados", a, len(s))
    reg("trazabilidad", "archivo_origen", "columna con la ruta del archivo de bronze", len(s), len(s))

    pct_nulos = {c: round(100.0 * float(s[c].isna().mean()), 2) for c in s.columns}
    controles = {
        "completitud_media_pct": completitud(s),
        "umbral_nulos_pct": UMBRAL_NULOS_PCT,
        "columnas_sobre_umbral": {c: p for c, p in pct_nulos.items() if p >= UMBRAL_NULOS_PCT},
    }
    ca, cm = roles["anio"], roles["mes"]
    if ca and ca in s.columns:
        anio = pd.to_numeric(s[ca], errors="coerce")
        controles["anio_distinto_de_fecha"] = int((anio.notna() & (s[F].dt.year != anio)).sum())
    if cm and cm in s.columns:
        mes_txt = s[cm].astype("string").str.strip().str.upper()
        mes = pd.to_numeric(mes_txt.map(MESES_ES), errors="coerce")
        mes = mes.fillna(pd.to_numeric(mes_txt, errors="coerce"))
        controles["mes_distinto_de_fecha"] = int((mes.notna() & (s[F].dt.month != mes)).sum())
    if cc and cc in s.columns:
        controles["ubigeo_formato_invalido"] = int((s[cc].notna() & ~s[cc].str.match(r"^\d{6}$", na=False)).sum())
    fmin, fmax = s[F].min(), s[F].max()
    controles["oportunidad"] = {
        "fecha_min": str(fmin.date()), "fecha_max": str(fmax.date()),
        "dias_desde_fecha_max": int((pd.Timestamp(datetime.now(ZoneInfo(TZ)).date()) - fmax).days),
    }
    despues = {"filas": len(s), "completitud_media_pct": completitud(s)}

    stem = basename(clave)[:-4]
    carpeta = f"silver/{INDECI}/{FECHA}"
    subir_bytes(f"{carpeta}/{stem}.xlsx", a_xlsx(s, {F: "DD-MM-YYYY"}), XLSX)
    subir_json(f"{carpeta}/{stem}_calidad.json", {
        "archivo_origen": clave, "fecha_proceso": FECHA,
        "antes": antes, "despues": despues, "controles": controles, "reglas_aplicadas": reglas,
    })
    cols_silver = list(s.columns)
    subir_json(f"{carpeta}/{stem}.json", {
        "capa": "silver", "fuente": INDECI_FUENTE, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "archivo_origen": clave, "fecha_origen": f_bronze,
        "archivo_destino": f"{carpeta}/{stem}.xlsx",
        "archivo_calidad": f"{carpeta}/{stem}_calidad.json",
        "periodo": f"{INICIO} a {FIN}",
        "filas_entrada": n0, "filas_salida": len(s),
        "roles": {k: v for k, v in roles.items() if v},
        "columnas_impacto": impacto,
        "columnas_agregadas": [c for c in cols_silver if c not in cols_bronze and c not in nuevos],
        "columna_final": cols_silver[-1],
        "tipos_de_datos": {c: str(t) for c, t in s.dtypes.items()},
        "estado": "procesado",
    })
    print(f"silver indeci: {n0:,} -> {len(s):,} filas -> s3://{BUCKET}/{carpeta}/{stem}.xlsx")


def indeci_gold():
    f_silver, clave = ultimo_simple("silver", INDECI, ".xlsx")
    df = pd.read_excel(io.BytesIO(leer_bytes(clave)), sheet_name="datos")
    cols = list(df.columns)
    fcol = next((c for c in cols if pd.api.types.is_datetime64_any_dtype(df[c])), None)
    if fcol is None:
        raise ValueError("No hay columna de fecha en el silver de INDECI")
    c_dpto = buscar(df, "dpto", "departamento")
    c_prov = buscar(df, "prov", "provincia")
    mapa = {
        "fallecidos": buscar(df, "fallecidos"),
        "desaparecidos": buscar(df, "desaparecidos"),
        "heridos": buscar(df, "lesionados", "heridos"),
        "damnificados": buscar(df, "damnificados"),
        "afectados": buscar(df, "afectados"),
        "viviendas_destruidas": buscar(df, "viviendas_destruidas"),
        "viviendas_afectadas": buscar(df, "viviendas_afectadas"),
        "centros_educativos_destruidos": buscar(df, "centros_educativos_destruidos"),
        "centros_salud_destruidos": buscar(df, "centros_salud_destruidos"),
        "has_cultivo_destruido": buscar(df, "has_cultivo_destruido"),
        "costo_de_la_ayuda": buscar(df, "costo_de_la_ayuda"),
    }
    g = pd.DataFrame({"fecha": df[fcol], "departamento": df[c_dpto], "provincia": df[c_prov]})
    for alias, col in mapa.items():
        g[alias] = pd.to_numeric(df[col], errors="coerce").fillna(0).clip(lower=0) if col else 0.0
    g["indice_urgencia"] = sum(g[a] * p for a, p in PESOS.items())

    cols_sum = list(PESOS.keys()) + ["costo_de_la_ayuda"]
    res = (g.groupby(LLAVES, dropna=False)
           .agg(n_emergencias=("indice_urgencia", "size"), indice_urgencia=("indice_urgencia", "sum"),
                **{a: (a, "sum") for a in cols_sum})
           .reset_index().sort_values(["fecha", "departamento", "provincia"]))
    res["indice_urgencia"] = res["indice_urgencia"].round(2)
    for c in [c for c in res.columns if c not in LLAVES]:
        res[c] = pd.to_numeric(res[c], errors="coerce").fillna(0)
    anios = res["fecha"].dt.year
    etiqueta = str(anios.min()) if anios.min() == anios.max() else f"{anios.min()}_{anios.max()}"
    res["fecha"] = res["fecha"].dt.strftime("%d-%m-%Y")
    destino = f"gold/{INDECI}/{FECHA}/indeci_{etiqueta}_{FECHA}.csv"
    subir(destino, res.to_csv(sep=SEP, index=False, na_rep="0"))
    print(f"gold indeci: {len(res)} filas (fecha, departamento, provincia) -> s3://{BUCKET}/{destino}")


# ========================== INGEMMET ==========================
# Inventario de Peligros Geológicos (INGEMMET). Alcance: nacional, periodo INICIO a FIN.

def ingemmet_fuente():
    clave = f"{CAPA_FUENTE}/{INGEMMET}/{FECHA}/ingemmet_peligros.csv"
    if existe(clave):
        print(f"Ya existe en S3, no se descarga: {clave}")
        return
    r = requests.get(IGP_API, params={"id": INGEMMET_DATASET_ID}, headers=UA, timeout=30)
    r.raise_for_status()
    res = r.json()["result"]
    dataset = res[0] if isinstance(res, list) else res
    csvs = [x for x in dataset["resources"]
            if x.get("url") and (x["url"].lower().endswith(".csv")
                                 or x.get("format", "").lower().lstrip(".") == "csv")]
    if not csvs:
        raise RuntimeError("El dataset de INGEMMET no tiene recursos CSV con URL")
    recurso = csvs[0]
    print("Descargando:", recurso["url"])
    archivo = requests.get(recurso["url"], headers=UA, timeout=300)
    archivo.raise_for_status()

    filas = list(csv.reader(io.StringIO(archivo.content.decode("latin-1")), delimiter=";"))
    encabezado = filas[0]
    subir_bytes(clave, archivo.content, "text/csv")
    subir_json(clave[:-4] + ".json", {
        "fuente": INGEMMET_FUENTE,
        "dataset": dataset.get("title"),
        "url_origen": recurso["url"],
        "fecha_modificacion_fuente": dataset.get("metadata_modified"),
        "fecha_hora_extraccion": datetime.now(ZoneInfo(TZ)).isoformat(),
        "cantidad_registros": len(filas) - 1,
        "cantidad_columnas": len(encabezado),
        "encabezado_columnas": encabezado,
        "formato": "csv", "origen_sep": ";", "origen_encoding": "latin-1",
        "bytes": len(archivo.content),
        "sha256": hashlib.sha256(archivo.content).hexdigest(),
        "ubicacion": f"s3://{BUCKET}/{clave}",
        "estado": "extraido",
    })
    print(f"fuente ingemmet: {len(filas) - 1:,} registros -> s3://{BUCKET}/{clave}")


def ingemmet_bronze():
    f_fuente, clave = ultimo_simple(CAPA_FUENTE, INGEMMET, ".csv")
    origen = leer_bytes(clave)
    ficha_fuente = json.loads(leer(clave[:-4] + ".json"))
    # dtype=str: se conserva todo tal cual viene (ubigeos, códigos, fechas)
    df = pd.read_csv(io.BytesIO(origen), sep=";", encoding="latin-1", low_memory=False, dtype=str)
    datos = df.to_csv(index=False, sep=SEP).encode("utf-8")

    nombre = f"ingemmet_peligros_{FECHA}"
    destino = f"bronze/{INGEMMET}/{FECHA}/{nombre}.csv"
    subir_bytes(destino, datos, "text/csv; charset=utf-8")
    subir_json(f"bronze/{INGEMMET}/{FECHA}/{nombre}.json", {
        "capa": "bronze", "fuente": INGEMMET_FUENTE, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "archivo_origen": clave, "fecha_origen": f_fuente,
        "fecha_hora_extraccion": ficha_fuente.get("fecha_hora_extraccion"),
        "fecha_modificacion_fuente": ficha_fuente.get("fecha_modificacion_fuente"),
        "url_origen": ficha_fuente.get("url_origen"),
        "archivo_destino": destino,
        "origen_sep": ";", "origen_encoding": "latin-1", "encoding": "utf-8", "sep": SEP,
        "filas_datos": int(len(df)), "n_columnas": int(len(df.columns)),
        "columnas": list(df.columns),
        "columna_final": df.columns[-1],
        "nulos_por_columna": {c: int(df[c].isna().sum()) for c in df.columns},
        "bytes_origen": len(origen), "bytes_bronze": len(datos),
        "sha256": hashlib.sha256(datos).hexdigest(),
        "estado": "copiado",
    })
    print(f"bronze ingemmet: {len(df):,} filas -> s3://{BUCKET}/{destino}")


REGLAS_ING = {
    "R1": "Completitud: código de inventario (COD_INV) no debe ser nulo",
    "R2": "Validez ubigeo: primeros 6 dígitos del COD_INV deben ser numéricos",
    "R3": "Filtro temporal: registros dentro del periodo INICIO a FIN",
    "R4": "Rango físico UTM: Este (100k-900k) y Norte (>8M)",
    "R5": "Unicidad: una sola fila por código de inventario (COD_INV)",
    "R6": "Validez de fecha: formato convertible a fecha",
    "R7": "Sin nulos: columnas finales de silver sin celdas vacías",
    "R8": "Altitud con dato (> 0 m): si no hay dato se marca altitud_valida = False (no se inventa)",
    "R9": f"Fallecidos atípicos (> {UMBRAL_FALLECIDOS_ATIPICO} en un solo registro): se marca dano_atipico = True",
}
REGLAS_MARCADAS_ING = {"R8", "R9"}  # no se eliminan ni se corrigen: se marcan en una columna


def calidad_ingemmet(df):
    cod = df["COD_INV"].astype(str).str.strip()
    f = parse_fecha(df["FECHA"])
    este = pd.to_numeric(df["ESTE"], errors="coerce")
    norte = pd.to_numeric(df["NORTE"], errors="coerce")
    alt = pd.to_numeric(df["ALTURA"], errors="coerce")
    muertos = pd.to_numeric(df["DA_MUERTO"], errors="coerce")
    return {
        "R1": int(df["COD_INV"].isna().sum()),
        "R2": int((~cod.str.zfill(11).str[:6].str.isdigit()).sum()),
        "R3": int((~f.between(pd.Timestamp(INICIO), pd.Timestamp(FIN))).sum()),
        "R4": int(((este < 100000) | (este > 900000) | (norte < 8000000)).sum()),
        "R5": int(cod.duplicated().sum()),
        "R6": int(f.isna().sum()),
        "R7": int(df.isna().sum().sum()),
        "R8": int((alt.isna() | (alt <= 0)).sum()),
        "R9": int((muertos > UMBRAL_FALLECIDOS_ATIPICO).sum()),
    }


def ingemmet_silver():
    f_bronze, clave = ultimo_simple("bronze", INGEMMET, ".csv")
    df_raw = pd.read_csv(io.BytesIO(leer_bytes(clave)), sep=SEP, dtype=str, low_memory=False)
    n0 = len(df_raw)
    cols_bronze = list(df_raw.columns)
    ev_antes = calidad_ingemmet(df_raw)
    acciones = []

    df = df_raw.copy()

    # A. Fecha (acepta varios formatos) y filtro del periodo
    df["fecha_dt"] = parse_fecha(df["FECHA"])
    a = len(df)
    df = df[df["fecha_dt"].between(pd.Timestamp(INICIO), pd.Timestamp(FIN))].copy()
    acciones.append((f"Filtrar periodo {INICIO} a {FIN} (incluye fechas inválidas)", a - len(df)))
    if df.empty:
        raise ValueError(f"INGEMMET no tiene peligros entre {INICIO} y {FIN}")
    df["fecha_registro"] = df["fecha_dt"].dt.normalize()
    df["anio_registro"] = df["fecha_dt"].dt.year

    # B. Código nulo y duplicados
    a = len(df)
    df = df[df["COD_INV"].notna()]
    acciones.append(("Eliminar filas sin COD_INV", a - len(df)))
    dups = int(df["COD_INV"].duplicated().sum())
    df = df.drop_duplicates(subset=["COD_INV"], keep="first")
    acciones.append(("Eliminar duplicados por COD_INV", dups))

    # C. Ubigeo de 6 dígitos
    df["ubigeo_distrito"] = df["COD_INV"].astype(str).str.strip().str.zfill(11).str[:6]
    acciones.append(("Derivar ubigeo_distrito (6 dígitos) desde COD_INV", 0))

    # D. Geografía normalizada igual que las otras fuentes (MAYÚSCULAS, sin tildes)
    n_geo = 0
    for c in ["NM_DPTO", "NM_PROV", "NM_DIST"]:
        if c in df.columns:
            df[c] = df[c].map(norm_txt).replace("", pd.NA)
            n_geo += int(df[c].isna().sum())
            df[c] = df[c].fillna("NO ESPECIFICADO")
    acciones.append(("Normalizar geografía (mayúsculas, sin tildes); nulos a 'NO ESPECIFICADO'", n_geo))

    # E. Métricas de daño: nulos -> 0, negativos -> 0
    total_num = 0
    for c in ["DA_VIV_AFECTADA", "DA_VIV_DESTRUIDA", "DA_DAMNIFICADO", "DA_MUERTO"]:
        if c in df.columns:
            v = pd.to_numeric(df[c], errors="coerce")
            total_num += int(v.isna().sum())
            df[c] = v.fillna(0).clip(lower=0).astype(int)
    acciones.append(("Imputar nulos en métricas numéricas de daño a 0", total_num))

    # F. Textos: nulos/vacíos -> 'NO ESPECIFICADO'
    total_txt = 0
    for c in ["NOM_CUENCA", "TIP_PELIGRO", "NOM_PELIGRO", "DENOMINACION", "DA_RIESGO_ESTIMADO", "EFECTUADO_POR"]:
        if c in df.columns:
            t = df[c].astype("string").str.strip()
            vac = (t.isna() | t.isin(["", "nan", "None", "NaN"])).fillna(True)
            total_txt += int(vac.sum())
            df[c] = t.mask(vac, "NO ESPECIFICADO").astype(object)
    acciones.append(("Imputar nulos en variables de texto a 'NO ESPECIFICADO'", total_txt))

    # G. Coordenadas y altitud. La altitud 0/nula NO es un dato real: se marca, no se inventa.
    for c in ["ESTE", "NORTE"]:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)
    alt = pd.to_numeric(df["ALTURA"], errors="coerce")
    df["altitud_valida"] = alt.gt(0).fillna(False)
    df["ALTURA"] = alt.fillna(0.0)
    acciones.append(("Marcar altitud sin dato (altitud_valida = False)", int((~df["altitud_valida"]).sum())))

    # G2. Fallecidos atípicos: se marcan para revisión, no se borran
    df["dano_atipico"] = df["DA_MUERTO"] > UMBRAL_FALLECIDOS_ATIPICO
    acciones.append((f"Marcar registros con más de {UMBRAL_FALLECIDOS_ATIPICO} fallecidos (dano_atipico = True)",
                     int(df["dano_atipico"].sum())))

    # H. Nombres finales
    columnas_map = {
        "COD_INV": "codigo_inventario", "ubigeo_distrito": "ubigeo_distrito",
        "NM_DPTO": "departamento", "NM_PROV": "provincia", "NM_DIST": "distrito",
        "NOM_CUENCA": "cuenca", "TIP_PELIGRO": "tipo_peligro", "NOM_PELIGRO": "nombre_peligro",
        "DENOMINACION": "denominacion", "ESTE": "coordenada_este", "NORTE": "coordenada_norte",
        "ZONA": "zona_utm", "ALTURA": "altitud_m", "altitud_valida": "altitud_valida",
        "fecha_registro": "fecha_registro", "anio_registro": "anio_registro",
        "DA_VIV_AFECTADA": "viviendas_afectadas", "DA_VIV_DESTRUIDA": "viviendas_destruidas",
        "DA_DAMNIFICADO": "personas_damnificadas", "DA_MUERTO": "personas_fallecidas",
        "dano_atipico": "dano_atipico",
        "DA_RIESGO_ESTIMADO": "nivel_riesgo_estimado", "EFECTUADO_POR": "evaluador_tecnico",
    }
    s = df[[c for c in columnas_map if c in df.columns]].rename(columns=columnas_map)
    s["archivo_origen"] = clave
    s = s.reset_index(drop=True)

    # Validación antes vs después
    ev_despues = calidad_ingemmet(df)
    ev_despues["R7"] = int(s.isna().sum().sum())
    reglas = []
    for k in REGLAS_ING:
        if ev_despues[k] == 0:
            estado = "CUMPLE"
        else:
            estado = "MARCADA" if k in REGLAS_MARCADAS_ING else "REVISAR"
        reglas.append({"regla": k, "descripcion": REGLAS_ING[k],
                       "incumplimientos_antes": ev_antes[k], "incumplimientos_despues": ev_despues[k],
                       "estado_final": estado})
    print(pd.DataFrame(acciones, columns=["Acción de limpieza", "Filas/valores afectados"]).to_string(index=False))
    print(pd.DataFrame(reglas)[["regla", "incumplimientos_antes", "incumplimientos_despues", "estado_final"]].to_string(index=False))

    atipicos = s.loc[s["dano_atipico"], ["codigo_inventario", "departamento", "provincia", "distrito",
                                         "tipo_peligro", "personas_fallecidas"]].to_dict("records")
    if atipicos:
        print(f"ATENCIÓN: {len(atipicos)} registro(s) con fallecidos atípicos, revisar con la fuente: {atipicos}")

    stem = basename(clave)[:-4]
    carpeta = f"silver/{INGEMMET}/{FECHA}"
    subir_bytes(f"{carpeta}/{stem}.xlsx", a_xlsx(s, {"fecha_registro": "DD-MM-YYYY"}), XLSX)
    subir_json(f"{carpeta}/{stem}_calidad.json", {
        "archivo_origen": clave, "fecha_proceso": FECHA,
        "reglas_calidad": reglas,
        "acciones_aplicadas": [{"accion": x, "afectados": int(n)} for x, n in acciones],
        "registros_con_dano_atipico": atipicos,
    })
    subir_json(f"{carpeta}/{stem}.json", {
        "capa": "silver", "fuente": INGEMMET_FUENTE, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "archivo_origen": clave, "fecha_origen": f_bronze,
        "archivo_destino": f"{carpeta}/{stem}.xlsx",
        "archivo_calidad": f"{carpeta}/{stem}_calidad.json",
        "periodo": f"{INICIO} a {FIN}",
        "alcance_geografico": "Nacional",
        "filas_entrada": n0, "filas_salida": len(s),
        "departamentos": sorted(s["departamento"].unique().tolist()),
        "columnas_originales": cols_bronze,
        "columnas_silver": list(s.columns),
        "columna_final": s.columns[-1],
        "tipos_de_datos": {c: str(t) for c, t in s.dtypes.items()},
        "estado": "procesado",
    })
    print(f"silver ingemmet: {n0:,} -> {len(s):,} filas -> s3://{BUCKET}/{carpeta}/{stem}.xlsx")


DIC_ING = {
    "diario_provincial": [
        ("fecha", "texto DD-MM-AAAA", "Fecha de registro del peligro"),
        ("departamento", "texto", "Departamento"),
        ("provincia", "texto", "Provincia"),
        ("n_peligros", "entero", "Cantidad de peligros inventariados ese día en la provincia"),
        ("viviendas_afectadas / viviendas_destruidas", "entero", "Sumas de daños a viviendas"),
        ("personas_damnificadas / personas_fallecidas", "entero", "Sumas de daños a personas"),
    ],
    "hechos": [
        ("codigo_inventario", "texto", "Identificador único INGEMMET (COD_INV)"),
        ("ubigeo_distrito", "texto (6 dig)", "UBIGEO del distrito (primeros 6 dígitos de COD_INV)"),
        ("departamento / provincia / distrito", "texto", "Geografía normalizada"),
        ("cuenca", "texto", "Cuenca hidrográfica"),
        ("tipo_peligro / nombre_peligro / denominacion", "texto", "Clasificación del peligro"),
        ("coordenada_este / coordenada_norte / zona_utm", "numérico", "Ubicación UTM"),
        ("altitud_m", "numérico", "Altitud en metros; 0 significa sin dato"),
        ("altitud_valida", "booleano", "True si la altitud tiene dato real (> 0)"),
        ("fecha_registro", "texto DD-MM-AAAA", "Fecha de registro"),
        ("anio_registro", "entero", "Año del registro"),
        ("viviendas_afectadas / viviendas_destruidas", "entero", "Daños a viviendas"),
        ("personas_damnificadas / personas_fallecidas", "entero", "Daños a personas"),
        ("dano_atipico", "booleano", f"True si el registro tiene más de {UMBRAL_FALLECIDOS_ATIPICO} fallecidos (revisar)"),
        ("nivel_riesgo_estimado / evaluador_tecnico", "texto", "Riesgo asignado y evaluador"),
        ("nivel_impacto", "texto", "Regla derivada: ALTO_DESTRUCCION, MEDIO_AFECTACION o SIN_DANO_DIRECTO_REGISTRADO"),
    ],
    "resumen_distrital": [
        ("ubigeo_distrito, departamento, provincia, distrito", "texto", "Llaves distritales"),
        ("total_peligros", "entero", "Total de peligros inventariados"),
        ("viviendas_afectadas / viviendas_destruidas", "entero", "Sumas de daños a viviendas"),
        ("personas_damnificadas / personas_fallecidas", "entero", "Sumas de daños a personas"),
    ],
    "por_tipo_peligro": [
        ("tipo_peligro, nombre_peligro", "texto", "Clasificación del peligro"),
        ("total_eventos", "entero", "Total de eventos"),
        ("viviendas_afectadas / viviendas_destruidas", "entero", "Sumas de daños a viviendas"),
    ],
    "anual_distrital": [
        ("ubigeo_distrito, departamento, distrito, anio_registro", "texto/entero", "Llaves distrito + año"),
        ("total_peligros", "entero", "Peligros del año"),
        ("viviendas_afectadas / viviendas_destruidas", "entero", "Sumas del año"),
    ],
}


def ingemmet_gold():
    f_silver, clave = ultimo_simple("silver", INGEMMET, ".xlsx")
    # dtype=str: si no, Excel devuelve ubigeos y códigos como números y se pierden los ceros
    h = pd.read_excel(io.BytesIO(leer_bytes(clave)), sheet_name="datos",
                      dtype={"codigo_inventario": str, "ubigeo_distrito": str, "zona_utm": str})
    h = h.drop(columns=["archivo_origen"])

    h["nivel_impacto"] = "SIN_DANO_DIRECTO_REGISTRADO"
    h.loc[h["viviendas_afectadas"] > 0, "nivel_impacto"] = "MEDIO_AFECTACION"
    h.loc[h["viviendas_destruidas"] > 0, "nivel_impacto"] = "ALTO_DESTRUCCION"

    danos = dict(viviendas_afectadas=("viviendas_afectadas", "sum"),
                 viviendas_destruidas=("viviendas_destruidas", "sum"),
                 personas_damnificadas=("personas_damnificadas", "sum"),
                 personas_fallecidas=("personas_fallecidas", "sum"))

    diario = (h.groupby(["fecha_registro", "departamento", "provincia"], as_index=False)
              .agg(n_peligros=("codigo_inventario", "count"), **danos)
              .rename(columns={"fecha_registro": "fecha"})
              .sort_values(["fecha", "departamento", "provincia"]))
    distrital = (h.groupby(["ubigeo_distrito", "departamento", "provincia", "distrito"], as_index=False)
                 .agg(total_peligros=("codigo_inventario", "count"), **danos))
    tipo = (h.groupby(["tipo_peligro", "nombre_peligro"], as_index=False)
            .agg(total_eventos=("codigo_inventario", "count"),
                 viviendas_afectadas=("viviendas_afectadas", "sum"),
                 viviendas_destruidas=("viviendas_destruidas", "sum"))
            .sort_values("total_eventos", ascending=False))
    anual = (h.groupby(["ubigeo_distrito", "departamento", "distrito", "anio_registro"], as_index=False)
             .agg(total_peligros=("codigo_inventario", "count"),
                  viviendas_afectadas=("viviendas_afectadas", "sum"),
                  viviendas_destruidas=("viviendas_destruidas", "sum")))

    anios = h["fecha_registro"].dt.year
    etiqueta = str(anios.min()) if anios.min() == anios.max() else f"{anios.min()}_{anios.max()}"
    diario["fecha"] = diario["fecha"].dt.strftime("%d-%m-%Y")
    h["fecha_registro"] = h["fecha_registro"].dt.strftime("%d-%m-%Y")

    # "hechos" es el que lee la integración (detalle por peligro con tipo, riesgo, cuenca y altitud)
    datasets = {"diario_provincial": diario, "hechos": h, "resumen_distrital": distrital,
                "por_tipo_peligro": tipo, "anual_distrital": anual}
    carpeta = f"gold/{INGEMMET}/{FECHA}"
    meta = []
    for nombre, frame in datasets.items():
        destino = f"{carpeta}/ingemmet_{nombre}_{etiqueta}_{FECHA}.csv"
        subir(destino, frame.to_csv(sep=SEP, index=False, na_rep=""))
        meta.append({"dataset": nombre, "registros": int(len(frame)),
                     "columnas": list(frame.columns), "ubicacion": f"s3://{BUCKET}/{destino}"})
        print(f"gold ingemmet/{nombre}: {len(frame):,} filas -> s3://{BUCKET}/{destino}")

    dic = pd.DataFrame([{"dataset": n, "columna": c, "tipo": t, "descripcion": d}
                        for n, filas in DIC_ING.items() for c, t, d in filas])
    k_dic = f"{carpeta}/ingemmet_diccionario_de_datos.csv"
    subir_bytes(k_dic, dic.to_csv(sep=SEP, index=False).encode("utf-8-sig"), "text/csv; charset=utf-8")
    subir_json(f"{carpeta}/ingemmet_gold_{FECHA}.json", {
        "capa": "gold", "fuente": INGEMMET_FUENTE, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "origen_silver": clave, "periodo": f"{INICIO} a {FIN}",
        "dataset_para_integracion": "hechos (detalle por peligro; la integración arma el perfil por provincia)",
        "datasets_gold": meta, "diccionario_datos": f"s3://{BUCKET}/{k_dic}",
        "estado": "procesado",
    })


# ======================== INTEGRACIÓN ========================
# IGP es la tabla base (left join). INDECI y Open-Meteo se cruzan por fecha + departamento + provincia.
# INGEMMET se integra como perfil geológico de la provincia (tipo de peligro, riesgo, cuenca,
# altitud, cantidades), sin repetir los campos de daños que ya trae INDECI.

def con_llaves(df, c_fecha, c_dpto, c_prov):
    out = df.copy()
    out["fecha"] = pd.to_datetime(out[c_fecha], format="%d-%m-%Y")
    out["departamento"] = out[c_dpto].map(norm_txt)
    out["provincia"] = out[c_prov].map(norm_txt)
    return out


def leer_gold(fuente, contiene=""):
    f, clave = ultimo_simple("gold", fuente, ".csv", contiene)
    print(f"gold {fuente}: {clave}")
    return clave, pd.read_csv(io.BytesIO(leer_bytes(clave)), sep=SEP, low_memory=False)


def homologar_indeci(df):
    out = df.copy()
    out.loc[out["departamento"] == "LIMA PROVINCIAS", "departamento"] = "LIMA"
    for (d, p), (d2, p2) in ALIAS_ZONAS.items():
        m = (out["departamento"] == d) & (out["provincia"] == p)
        out.loc[m, "departamento"] = d2
        out.loc[m, "provincia"] = p2
    return out


def a_xlsx_integracion(df):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="xlsxwriter", date_format="DD-MM-YYYY", datetime_format="DD-MM-YYYY") as w:
        df.to_excel(w, index=False, sheet_name="datos")
        wb, ws = w.book, w.sheets["datos"]
        izq = wb.add_format({"align": "left"})
        der = wb.add_format({"align": "right"})
        ws.freeze_panes(1, 3)
        for j, c in enumerate(df.columns):
            s = df[c]
            derecha = ((pd.api.types.is_numeric_dtype(s) and not pd.api.types.is_bool_dtype(s))
                       or pd.api.types.is_datetime64_any_dtype(s))
            largo = max([len(c)] + s.astype(str).str.len().head(200).tolist())
            ws.set_column(j, j, min(largo + 2, 40), der if derecha else izq)
    return buf.getvalue()


VENTANA_INGEMMET_DIAS = 30  # peligros INGEMMET en la misma provincia a ±N días del sismo

NIVELES_RIESGO = {"MUY ALTO": 4, "ALTO": 3, "MEDIO": 2, "BAJO": 1}
NIVELES_RIESGO_INV = {4: "MUY ALTO", 3: "ALTO", 2: "MEDIO", 1: "BAJO", 0: "NO ESPECIFICADO"}
CLASES_PELIGRO = ("deslizamiento", "erosion", "flujo", "caida", "inundacion", "otros")
SIN_REGISTRO = "SIN REGISTRO"


def clase_peligro(t):
    t = norm_txt(t)
    t = "" if pd.isna(t) else t
    for clave, nombre in (("DESLIZ", "deslizamiento"), ("EROSION", "erosion"), ("FLUJO", "flujo"),
                          ("CAIDA", "caida"), ("INUNDACION", "inundacion")):
        if clave in t:
            return nombre
    return "otros"


def perfil_ingemmet(ev):
    """Un perfil por (departamento, provincia) con campos que INDECI NO tiene."""
    ev = ev.copy()
    for c in ("tipo_peligro", "nombre_peligro", "cuenca"):
        ev[c] = ev[c].map(norm_txt).fillna("NO ESPECIFICADO")
    ev["_clase"] = ev["tipo_peligro"].map(clase_peligro)
    ev["_nivel"] = ev["nivel_riesgo_estimado"].map(norm_txt).map(NIVELES_RIESGO).fillna(0).astype(int)
    ev["_alto"] = (ev["_nivel"] >= 3).astype(int)
    # altitud 0 en INGEMMET significa "sin dato": solo se promedian altitudes válidas
    alt = pd.to_numeric(ev["altitud_m"], errors="coerce")
    if "altitud_valida" in ev.columns:
        valida = ev["altitud_valida"].astype(str).str.upper().eq("TRUE")
    else:
        valida = alt > 0
    ev["_alt"] = alt.where(valida & (alt > 0))

    moda = lambda s: s.value_counts().idxmax()
    base = ev.groupby(["departamento", "provincia"]).agg(
        ingemmet_n_peligros_provincia=("tipo_peligro", "size"),
        ingemmet_tipo_peligro_predominante=("tipo_peligro", moda),
        ingemmet_nombre_peligro_predominante=("nombre_peligro", moda),
        ingemmet_cuenca_principal=("cuenca", moda),
        ingemmet_n_riesgo_alto=("_alto", "sum"),
        _nivel_max=("_nivel", "max"),
        ingemmet_altitud_prom_m=("_alt", "mean"),
    )
    clases = (ev.groupby(["departamento", "provincia", "_clase"]).size()
              .unstack(fill_value=0)
              .reindex(columns=CLASES_PELIGRO, fill_value=0)
              .add_prefix("ingemmet_n_"))
    perfil = base.join(clases)
    perfil["ingemmet_nivel_riesgo_max"] = perfil.pop("_nivel_max").map(NIVELES_RIESGO_INV)
    perfil["ingemmet_altitud_prom_m"] = perfil["ingemmet_altitud_prom_m"].fillna(0).round(0)
    return perfil.reset_index()


# ---------------------- gold_integracion NUEVA ----------------------

def gold_integracion():
    k_om, om = leer_gold(FUENTE)
    k_igp, igp = leer_gold(IGP)
    k_ind, ind = leer_gold(INDECI)
    # Ahora se lee el detalle por evento (hechos), no el resumen diario
    k_ing, ing = leer_gold(INGEMMET, "_hechos_")
    # Reporte (noticias): se lee el gold agregado por fecha + departamento
    k_rep, rep = leer_gold(REPORTE, "_departamental_")
    k_rnac, rnac = leer_gold(REPORTE, "_nacional_")
    print(f"filas leídas -> open_meteo: {len(om)} | igp: {len(igp)} | indeci: {len(ind)} | "
          f"ingemmet (peligros): {len(ing)} | reporte (fecha-departamento): {len(rep)} | "
          f"reporte (nacional, días): {len(rnac)}")

    om = con_llaves(om, "fecha", "departamento", "provincia")
    om["_om"] = True

    ig = con_llaves(igp, "FECHA_UTC", "DEPARTAMENTO_CERCANO", "PROVINCIA_CERCANA")
    ig = (ig.groupby(LLAVES, dropna=False)
          .agg(igp_n_sismos=("ID", "count"), igp_magnitud_max=("MAGNITUD", "max"),
               igp_magnitud_prom=("MAGNITUD", "mean"), igp_profundidad_prom=("PROFUNDIDAD", "mean"))
          .round(2).reset_index())
    ig["_igp"] = True

    ind = homologar_indeci(con_llaves(ind, "fecha", "departamento", "provincia"))
    valores = [c for c in ind.columns if c not in LLAVES]
    ind = ind.groupby(LLAVES, dropna=False, as_index=False)[valores].sum()
    ind = ind.rename(columns={c: f"indeci_{c}" for c in valores})
    ind["_indeci"] = True

    # ---------- REPORTE: noticias (las noticias no traen provincia) ----------
    # (a) mismo día + departamento del sismo
    rep["fecha"] = pd.to_datetime(rep["fecha"], format="%d-%m-%Y")
    rep["departamento"] = rep["departamento"].map(norm_txt)
    rep = rep.groupby(["fecha", "departamento"], as_index=False).sum(numeric_only=True)
    valores_rep = [c for c in rep.columns if c not in ("fecha", "departamento")]

    # (b) ventana: mismo departamento, desde el día del sismo hasta N días después
    cruce_rep = ig[["fecha", "departamento"]].drop_duplicates().merge(
        rep.rename(columns={"fecha": "fecha_noticia"}), on="departamento")
    cruce_rep["dias"] = (cruce_rep["fecha_noticia"] - cruce_rep["fecha"]).dt.days
    cruce_rep = cruce_rep[cruce_rep["dias"].between(0, VENTANA_REPORTE_DIAS)]
    rep_ventana = (cruce_rep.groupby(["fecha", "departamento"], as_index=False)[valores_rep].sum()
                   .rename(columns={c: f"reporte_{c}_{VENTANA_REPORTE_DIAS}d" for c in valores_rep}))

    # (c) nacional: todas las noticias del día, con o sin departamento
    rnac["fecha"] = pd.to_datetime(rnac["fecha"], format="%d-%m-%Y")
    rnac = rnac.groupby("fecha", as_index=False).sum(numeric_only=True)
    rnac = rnac.rename(columns={c: f"reporte_nac_{c}" for c in rnac.columns if c != "fecha"})

    rep = rep.rename(columns={c: f"reporte_{c}" for c in valores_rep})
    cols_rep = [f"reporte_{c}" for c in valores_rep]
    cols_rep_vent = [f"reporte_{c}_{VENTANA_REPORTE_DIAS}d" for c in valores_rep]
    cols_rep_nac = [c for c in rnac.columns if c != "fecha"]

    # ---------- INGEMMET: perfil de la provincia (campos que INDECI no tiene) ----------
    ev = homologar_indeci(con_llaves(ing, "fecha_registro", "departamento", "provincia"))
    perfil = perfil_ingemmet(ev)

    n_dep = ev.groupby("departamento").size().rename("ingemmet_n_peligros_departamento").reset_index()

    # peligros de la misma provincia cerca de la fecha del sismo (mismo día y ventana ±N días)
    cruce = ig[LLAVES].merge(
        ev[["departamento", "provincia", "fecha"]].rename(columns={"fecha": "fecha_peligro"}),
        on=["departamento", "provincia"])
    cruce["dias"] = (cruce["fecha"] - cruce["fecha_peligro"]).abs().dt.days
    col_ventana = f"ingemmet_n_peligros_{VENTANA_INGEMMET_DIAS}d"
    mismo_dia = (cruce[cruce["dias"] == 0].groupby(LLAVES).size()
                 .rename("ingemmet_n_peligros_mismo_dia").reset_index())
    ventana = (cruce[cruce["dias"] <= VENTANA_INGEMMET_DIAS].groupby(LLAVES).size()
               .rename(col_ventana).reset_index())

    pares = lambda d: set(d[["departamento", "provincia"]].dropna().drop_duplicates().itertuples(index=False, name=None))
    zonas_om = pares(om)
    sin_cruce = {
        "igp": sorted(f"{d} / {p}" for d, p in pares(ig) - zonas_om),
        "indeci": sorted(f"{d} / {p}" for d, p in pares(ind) - zonas_om),
        "ingemmet": sorted(f"{d} / {p}" for d, p in pares(perfil) - zonas_om),
    }

    # IGP a la izquierda: manda la base. Si no hay sismo en IGP, la fila no aparece.
    m = (ig.merge(ind, on=LLAVES, how=TIPO_CRUCE)
           .merge(om, on=LLAVES, how=TIPO_CRUCE)
           .merge(rep, on=["fecha", "departamento"], how=TIPO_CRUCE)
           .merge(rep_ventana, on=["fecha", "departamento"], how=TIPO_CRUCE)
           .merge(rnac, on="fecha", how=TIPO_CRUCE)
           .merge(perfil, on=["departamento", "provincia"], how=TIPO_CRUCE)
           .merge(n_dep, on="departamento", how=TIPO_CRUCE)
           .merge(mismo_dia, on=LLAVES, how=TIPO_CRUCE)
           .merge(ventana, on=LLAVES, how=TIPO_CRUCE))
    assert m["_igp"].eq(True).all(), "Hay filas sin IGP: revisa que TIPO_CRUCE sea 'left'"
    assert len(m) == len(ig), "El cruce cambió el número de filas de IGP: hay llaves repetidas"

    cols_ing = ["ingemmet_n_peligros_mismo_dia", col_ventana,
                "ingemmet_n_peligros_provincia", "ingemmet_n_peligros_departamento",
                "ingemmet_tipo_peligro_predominante", "ingemmet_nombre_peligro_predominante",
                "ingemmet_n_deslizamiento", "ingemmet_n_erosion", "ingemmet_n_flujo",
                "ingemmet_n_caida", "ingemmet_n_inundacion", "ingemmet_n_otros",
                "ingemmet_nivel_riesgo_max", "ingemmet_n_riesgo_alto",
                "ingemmet_cuenca_principal", "ingemmet_altitud_prom_m"]
    cols_txt = ["ingemmet_tipo_peligro_predominante", "ingemmet_nombre_peligro_predominante",
                "ingemmet_nivel_riesgo_max", "ingemmet_cuenca_principal"]
    cols_num = [c for c in cols_ing if c not in cols_txt]

    # tiene_ingemmet = la provincia tiene al menos un peligro inventariado en el periodo
    m["tiene_ingemmet"] = m["ingemmet_n_peligros_provincia"].notna()
    m[cols_num] = m[cols_num].fillna(0)
    for c in cols_num:
        if c != "ingemmet_altitud_prom_m":
            m[c] = m[c].astype(int)
    m[cols_txt] = m[cols_txt].fillna(SIN_REGISTRO)

    m["tiene_igp"] = m["_igp"].eq(True)
    m["tiene_indeci"] = m["_indeci"].eq(True)
    m["tiene_open_meteo"] = m["_om"].eq(True)

    # tiene_reporte = ese día hubo al menos una noticia que menciona el departamento del sismo
    todas_rep = cols_rep + cols_rep_vent + cols_rep_nac
    m[todas_rep] = m[todas_rep].fillna(0).astype(int)
    m["tiene_reporte"] = m["reporte_n_noticias"] > 0

    flags = ["tiene_igp", "tiene_indeci", "tiene_ingemmet", "tiene_reporte", "tiene_open_meteo"]
    cols_igp = [c for c in ig.columns if c not in LLAVES and not c.startswith("_")]
    cols_ind = [c for c in ind.columns if c not in LLAVES and not c.startswith("_")]
    cols_om = [c for c in GOLD_COLS if c not in LLAVES]

    m = m.sort_values(["departamento", "provincia", "fecha"])
    anios = m["fecha"].dt.year
    etiqueta = str(anios.min()) if anios.min() == anios.max() else f"{anios.min()}_{anios.max()}"
    m = m[LLAVES + flags + cols_igp + cols_ind + cols_ing + todas_rep + cols_om].reset_index(drop=True)

    # INDECI: cualquier celda vacía/NaN en columnas indeci_* → 0 (también filas sin match en left join)
    for c in [c for c in m.columns if c.startswith("indeci_")]:
        m[c] = pd.to_numeric(m[c], errors="coerce").fillna(0)

    datos = a_xlsx_integracion(m)
    destino = f"gold/integracion_de_datos/{FECHA}/integracion_{etiqueta}_{FECHA}.xlsx"
    subir_bytes(destino, datos, XLSX)
    destino_csv = destino[:-5] + ".csv"
    m_csv = m.copy()
    m_csv["fecha"] = m_csv["fecha"].dt.strftime("%d-%m-%Y")
    # Forzar 0 también al exportar (por si queda algún NaN)
    for c in [c for c in m_csv.columns if c.startswith("indeci_")]:
        m_csv[c] = pd.to_numeric(m_csv[c], errors="coerce").fillna(0)
    subir(destino_csv, m_csv.to_csv(sep=SEP, index=False, na_rep="0"))

    n_fuentes = m[flags].sum(axis=1)
    subir_json(destino[:-5] + ".json", {
        "capa": "gold", "dataset": "integracion_de_datos", "fecha_proceso": FECHA, "zona_horaria": TZ,
        "formato": "xlsx + csv", "hoja": "datos",
        "fuentes": {"igp": k_igp, "indeci": k_ind, "ingemmet": k_ing, "reporte": k_rep, "reporte_nacional": k_rnac, "open_meteo": k_om},
        "llaves": LLAVES, "tipo_cruce": TIPO_CRUCE,
        "base": "igp (INDECI y Open-Meteo se cruzan por fecha + departamento + provincia)",
        "ingemmet_modo": f"perfil de la provincia (departamento + provincia) leído del gold 'hechos'; "
                         f"ventana de ±{VENTANA_INGEMMET_DIAS} días respecto al sismo; "
                         f"sin registro = 0 en conteos y '{SIN_REGISTRO}' en texto; "
                         "no repite los campos de daños que ya trae INDECI",
        "reporte_modo": "tres bloques: (1) reporte_*: noticias del mismo día y departamento (gold 'departamental'); "
                        f"(2) reporte_*_{VENTANA_REPORTE_DIAS}d: noticias del mismo departamento desde el día del sismo hasta "
                        f"{VENTANA_REPORTE_DIAS} días después; (3) reporte_nac_*: todas las noticias del día en el país (gold 'nacional'). "
                        "Como las noticias no traen provincia, el conteo se repite en todas las provincias del departamento; sin noticias = 0",
        "orden_columnas": "llaves, banderas, IGP, INDECI, INGEMMET, Reporte (día, ventana, nacional), Open-Meteo",
        "normalizacion_llaves": "departamento y provincia en MAYÚSCULAS, sin tildes y sin espacios repetidos",
        "homologacion_zonas_indeci_e_ingemmet": {f"{d} / {p}": f"{d2} / {p2}" for (d, p), (d2, p2) in ALIAS_ZONAS.items()},
        "igp_agregacion": "por fecha UTC, departamento y provincia más cercana a cada epicentro",
        "filas": len(m), "columnas": list(m.columns),
        "filas_con_indeci": int(m["tiene_indeci"].sum()),
        "filas_con_reporte": int(m["tiene_reporte"].sum()),
        "filas_con_reporte_ventana": int((m[f"reporte_n_noticias_{VENTANA_REPORTE_DIAS}d"] > 0).sum()),
        "filas_con_reporte_nacional": int((m["reporte_nac_n_noticias"] > 0).sum()),
        "filas_con_ingemmet_en_provincia": int(m["tiene_ingemmet"].sum()),
        "filas_con_ingemmet_en_ventana": int((m[col_ventana] > 0).sum()),
        "filas_con_ingemmet_mismo_dia": int((m["ingemmet_n_peligros_mismo_dia"] > 0).sum()),
        "provincias_sin_cruce_con_open_meteo": sin_cruce,
        "sha256": hashlib.sha256(datos).hexdigest(),
        "ubicacion": f"s3://{BUCKET}/{destino}",
        "ubicacion_csv": f"s3://{BUCKET}/{destino_csv}",
    })
    print(f"\nintegración: {len(m)} filas (base IGP) | con INDECI: {int(m['tiene_indeci'].sum())} | "
          f"con noticias (reporte, mismo día): {int(m['tiene_reporte'].sum())} | "
          f"ventana {VENTANA_REPORTE_DIAS}d: {int((m[f'reporte_n_noticias_{VENTANA_REPORTE_DIAS}d'] > 0).sum())} | "
          f"nacional: {int((m['reporte_nac_n_noticias'] > 0).sum())}")
    print(f"INGEMMET -> provincia con peligros: {int(m['tiene_ingemmet'].sum())} filas | "
          f"a ±{VENTANA_INGEMMET_DIAS} días del sismo: {int((m[col_ventana] > 0).sum())} | "
          f"mismo día: {int((m['ingemmet_n_peligros_mismo_dia'] > 0).sum())}")
    print(f"provincias sin cruce -> IGP: {len(sin_cruce['igp'])} | INDECI: {len(sin_cruce['indeci'])} | "
          f"INGEMMET: {len(sin_cruce['ingemmet'])}", sin_cruce["ingemmet"])
    print(f"-> s3://{BUCKET}/{destino}")
    print(f"-> s3://{BUCKET}/{destino_csv}")


REPORTE = "reporte"
REPORTE_FUENTE = "Google News RSS (noticias públicas sobre peligros y emergencias en Perú)"
NOTICIAS_RSS = "https://news.google.com/rss/search"
NOTICIAS_TEMAS = ["sismo Perú", "huaico Perú", "inundación Perú", "deslizamiento Perú",
                  "emergencia INDECI Perú", "lluvias intensas Perú", "estado de emergencia Perú"]
NOTICIAS_PAUSA = 1.0
NOTICIAS_VENTANA = "semana"  # "mes" (menos consultas) o "semana" (menos noticias perdidas por el tope de ~100 de Google News)
# Consultas por departamento ("huaico Cusco"): recuperan noticias locales que la consulta nacional no trae.
# Se hacen por mes para no multiplicar las consultas (25 departamentos x temas x 12 meses).
NOTICIAS_POR_DEPARTAMENTO = True
NOTICIAS_TEMAS_DEPTO = ["huaico", "inundación", "deslizamiento"]
HILOS_NOTICIAS = 4  # consultas a Google News en paralelo; si empieza a dar 429, baja a 2
CACHE_NOTICIAS = "/tmp/cache_noticias"  # disco local del cluster: permite reanudar si se corta la descarga
FORZAR_REPORTE = False  # True = vuelve a descargar aunque ya exista el archivo de fuente de hoy
VENTANA_REPORTE_DIAS = 3  # noticias del mismo departamento desde el día del sismo hasta N días después
TEMAS_REPORTE = ["sismo", "huaico", "inundacion", "deslizamiento", "emergencia", "lluvia"]
DEPARTAMENTOS_PE = [
    "AMAZONAS", "ANCASH", "APURIMAC", "AREQUIPA", "AYACUCHO", "CAJAMARCA", "CALLAO",
    "CUSCO", "HUANCAVELICA", "HUANUCO", "ICA", "JUNIN", "LA LIBERTAD", "LAMBAYEQUE",
    "LIMA", "LORETO", "MADRE DE DIOS", "MOQUEGUA", "PASCO", "PIURA", "PUNO",
    "SAN MARTIN", "TACNA", "TUMBES", "UCAYALI",
]


def meses_rango(ini, fin):
    a = date(ini.year, ini.month, 1)
    while a <= fin:
        sig = (a.replace(day=28) + timedelta(days=4)).replace(day=1)
        yield max(a, ini), min(sig - timedelta(days=1), fin)
        a = sig


def semanas_rango(ini, fin):
    a = ini
    while a <= fin:
        b = min(a + timedelta(days=6), fin)
        yield a, b
        a = b + timedelta(days=1)


def limpiar_html(t):
    t = re.sub(r"<[^>]+>", " ", html.unescape(t or ""))
    return re.sub(r"\s+", " ", t).strip()


def detectar_departamento(texto):
    # devuelve el departamento que aparece PRIMERO en el texto (no el primero alfabético)
    t = f" {norm_txt(texto) or ''} "
    mejor = None
    for d in DEPARTAMENTOS_PE:
        m = re.search(rf"\b{re.escape(d)}\b", t)
        if m and (mejor is None or m.start() < mejor[0]):
            mejor = (m.start(), d)
    return mejor[1] if mejor else ""


def pedir_noticias(tema, a, b):
    q = f"{tema} after:{a.isoformat()} before:{(b + timedelta(days=1)).isoformat()}"
    params = {"q": q, "hl": "es-419", "gl": "PE", "ceid": "PE:es-419"}
    for intento in range(4):
        r = requests.get(NOTICIAS_RSS, params=params, headers=UA, timeout=30)
        if r.status_code == 200:
            return ET.fromstring(r.content)
        if r.status_code in (429, 500, 502, 503, 504):
            time.sleep(5 * (intento + 1))
            continue
        r.raise_for_status()
    raise RuntimeError(f"Google News no respondió para: {q}")


def consultas_noticias():
    """Lista de (texto de la consulta, departamento de la consulta, función de rangos)."""
    rangos = semanas_rango if NOTICIAS_VENTANA == "semana" else meses_rango
    for tema in NOTICIAS_TEMAS:
        yield tema, "", rangos
    if NOTICIAS_POR_DEPARTAMENTO:
        for dep in DEPARTAMENTOS_PE:
            for tema in NOTICIAS_TEMAS_DEPTO:
                yield f"{tema} {dep.title()}", dep, meses_rango


def noticias_de_consulta(tema, dep_consulta, a, b):
    """Una consulta (tema, rango) -> lista de noticias. Guarda el resultado en disco: si se corta, no se repite."""
    archivo = os.path.join(CACHE_NOTICIAS, hashlib.sha1(f"{tema}|{a}|{b}".encode("utf-8")).hexdigest() + ".json")
    if os.path.exists(archivo):
        with open(archivo, encoding="utf-8") as fh:
            return json.load(fh)
    raiz = pedir_noticias(tema, a, b)
    filas = []
    for it in raiz.iter("item"):
        titulo = limpiar_html(it.findtext("title"))
        url = (it.findtext("link") or "").strip()
        pub = it.findtext("pubDate")
        if not titulo or not url or not pub:
            continue
        try:
            dt = parsedate_to_datetime(pub).astimezone(ZoneInfo(TZ))
        except Exception:
            continue
        if not (INICIO <= dt.date() <= FIN):
            continue
        medio = limpiar_html(it.findtext("source"))
        resumen = limpiar_html(it.findtext("description"))
        if resumen == titulo or resumen.startswith(titulo):
            resumen = ""
        filas.append({
            "fecha": dt.strftime("%d-%m-%Y"),
            "hora_lima": dt.strftime("%H:%M:%S"),
            "titulo": titulo,
            "medio": medio,
            "url": url,
            "resumen": resumen,
            "tema_busqueda": tema,
            "departamento_consulta": dep_consulta,
            "departamento_mencionado": detectar_departamento(f"{titulo} {resumen}"),
        })
    os.makedirs(CACHE_NOTICIAS, exist_ok=True)
    with open(archivo, "w", encoding="utf-8") as fh:
        json.dump(filas, fh, ensure_ascii=False)
    time.sleep(NOTICIAS_PAUSA)
    return filas


def extraer_noticias():
    """Descarga en paralelo (no usa AWS). Cada consulta queda en la caché local CACHE_NOTICIAS."""
    tareas = [(tema, dep, a, b) for tema, dep, rangos in consultas_noticias() for a, b in rangos(INICIO, FIN)]
    en_cache = sum(os.path.exists(os.path.join(
        CACHE_NOTICIAS, hashlib.sha1(f"{t}|{a}|{b}".encode("utf-8")).hexdigest() + ".json"))
        for t, _, a, b in tareas)
    print(f"{len(tareas)} consultas ({en_cache} ya estaban en la caché local), {HILOS_NOTICIAS} en paralelo.")
    resultados, fallidas = {}, []
    with ThreadPoolExecutor(HILOS_NOTICIAS) as ex:
        futuros = {ex.submit(noticias_de_consulta, *t): i for i, t in enumerate(tareas)}
        for n, f in enumerate(as_completed(futuros), 1):
            i = futuros[f]
            try:
                resultados[i] = f.result()
            except Exception as e:
                tema, _, a, b = tareas[i]
                fallidas.append(f"{tema} {a} a {b}: {e!r}")
            if n % 50 == 0 or n == len(tareas):
                total = sum(len(v) for v in resultados.values())
                print(f"  {n}/{len(tareas)} consultas | {total} noticias | {len(fallidas)} fallidas")
    filas = []
    for i in sorted(resultados):
        for fila in resultados[i]:
            filas.append({**fila, "fecha_extraccion": FECHA})
    return filas, fallidas


def reporte_fuente():
    clave = f"{CAPA_FUENTE}/{REPORTE}/{FECHA}/reporte_noticias.csv"
    if existe(clave) and not FORZAR_REPORTE:
        print(f"Ya existe en S3, no se descarga: {clave} (usa FORZAR_REPORTE = True para volver a descargar)")
        return
    filas, fallidas = extraer_noticias()
    if not filas:
        raise RuntimeError(f"No se obtuvo ninguna noticia ({len(fallidas)} consultas fallidas)")
    df = pd.DataFrame(filas)
    n_bruto = len(df)
    df = df.drop_duplicates(subset=["url"]).drop_duplicates(subset=["titulo", "medio"])
    df = df.assign(_f=pd.to_datetime(df["fecha"], format="%d-%m-%Y")) \
           .sort_values(["_f", "hora_lima"]).drop(columns="_f").reset_index(drop=True)

    datos = df.to_csv(sep=SEP, index=False, na_rep="")
    subir(clave, datos)
    subir_json(clave[:-4] + ".json", {
        "fuente": REPORTE_FUENTE,
        "dataset": "noticias_peligros_peru",
        "url_origen": NOTICIAS_RSS,
        "temas": NOTICIAS_TEMAS,
        "temas_por_departamento": NOTICIAS_TEMAS_DEPTO if NOTICIAS_POR_DEPARTAMENTO else [],
        "ventana_consulta_nacional": NOTICIAS_VENTANA,
        "periodo": f"{INICIO} a {FIN}",
        "fecha_hora_extraccion": datetime.now(ZoneInfo(TZ)).isoformat(),
        "cantidad_registros": int(len(df)),
        "duplicados_eliminados": int(n_bruto - len(df)),
        "consultas_fallidas": fallidas,
        "columnas": list(df.columns),
        "formato": "csv", "sep": SEP, "encoding": "utf-8",
        "bytes": len(datos.encode("utf-8")),
        "sha256": hashlib.sha256(datos.encode("utf-8")).hexdigest(),
        "ubicacion": f"s3://{BUCKET}/{clave}",
        "estado": "extraido",
    })
    print(f"fuente reporte: {len(df):,} noticias -> s3://{BUCKET}/{clave}")


def clasificar_tema(t):
    t = norm_txt(t) or ""
    for clave, nombre in (("SISMO", "sismo"), ("HUAICO", "huaico"), ("INUNDACION", "inundacion"),
                          ("DESLIZ", "deslizamiento"), ("EMERGENCIA", "emergencia"),
                          ("LLUVIA", "lluvia")):
        if clave in t:
            return nombre
    return "otros"


def quitar_medio(titulo, medio):
    """Google News agrega ' - Medio' al final del título; el medio ya tiene su propia columna."""
    if medio:
        for sep in (" - ", " | "):
            suf = f"{sep}{medio}"
            if titulo.endswith(suf):
                return titulo[:-len(suf)].strip()
    return titulo


def calidad_reporte(df):
    vacio = df.replace("", pd.NA)
    return {
        "filas": len(df),
        "completitud_pct": completitud(vacio),
        "vacios_por_columna": {c: int(n) for c, n in vacio.isna().sum().items() if n},
        "urls_duplicadas": int(df["url"].duplicated().sum()),
    }


def reporte_bronze():
    f_fuente, clave = ultimo_simple(CAPA_FUENTE, REPORTE, ".csv")
    datos = leer_bytes(clave)
    ficha_fuente = json.loads(leer(clave[:-4] + ".json"))
    df = pd.read_csv(io.BytesIO(datos), sep=SEP, dtype=str, keep_default_na=False)
    nombre = f"reporte_noticias_{FECHA}"
    destino = f"bronze/{REPORTE}/{FECHA}/{nombre}.csv"
    subir_bytes(destino, datos, "text/csv; charset=utf-8")
    subir_json(f"bronze/{REPORTE}/{FECHA}/{nombre}.json", {
        "capa": "bronze", "fuente": REPORTE_FUENTE, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "archivo_origen": clave, "fecha_origen": f_fuente,
        "fecha_hora_extraccion": ficha_fuente.get("fecha_hora_extraccion"),
        "url_origen": ficha_fuente.get("url_origen"),
        "temas": ficha_fuente.get("temas"),
        "archivo_destino": destino,
        "filas_datos": int(len(df)), "n_columnas": int(len(df.columns)),
        "columnas": list(df.columns),
        "columna_final": df.columns[-1],
        "bytes": len(datos),
        "sha256": hashlib.sha256(datos).hexdigest(),
        "estado": "copiado",
    })
    print(f"bronze reporte: {len(df):,} filas -> s3://{BUCKET}/{destino}")


def reporte_silver():
    f_bronze, clave = ultimo_simple("bronze", REPORTE, ".csv")
    df = pd.read_csv(io.BytesIO(leer_bytes(clave)), sep=SEP, dtype=str, keep_default_na=False)
    n0 = len(df)
    cols_bronze = list(df.columns)
    antes = calidad_reporte(df)
    reglas = []

    def reg(accion, campo, detalle, motivo, fa, fd):
        reglas.append({"accion": accion, "campo": campo, "detalle": detalle, "motivo": motivo,
                       "filas_antes": fa, "filas_despues": fd})

    if "departamento_consulta" not in df.columns:
        df["departamento_consulta"] = ""
    for c in ["titulo", "medio", "url", "resumen"]:
        df[c] = df[c].str.replace(r"\s+", " ", regex=True).str.strip()
    reg("normalizar_texto", "titulo, medio, url, resumen", "espacios repetidos y bordes", "texto consistente", n0, n0)

    antes_t = df["titulo"].copy()
    df["titulo"] = [quitar_medio(t, m) for t, m in zip(df["titulo"], df["medio"])]
    reg("normalizar_texto", "titulo", f"{int((antes_t != df['titulo']).sum())} títulos sin el sufijo del medio",
        "el medio ya está en su propia columna y ensuciaba el deduplicado", n0, n0)

    df["fecha"] = pd.to_datetime(df["fecha"], format="%d-%m-%Y", errors="coerce")
    a = len(df)
    df = df[df["fecha"].notna()].copy()
    reg("eliminar_filas", "fecha", f"{a - len(df)} filas con fecha inválida", "no se pueden ubicar en el tiempo", a, len(df))
    a = len(df)
    df = df[df["fecha"].between(pd.Timestamp(INICIO), pd.Timestamp(FIN))].copy()
    reg("filtrar_filas", "fecha", f"{a - len(df)} filas fuera de {INICIO} a {FIN}", "solo el periodo solicitado", a, len(df))
    if df.empty:
        raise ValueError(f"No hay noticias entre {INICIO} y {FIN}")

    a = len(df)
    df = df[(df["titulo"] != "") & (df["url"] != "")].copy()
    reg("eliminar_filas", "titulo, url", f"{a - len(df)} filas sin título o sin URL", "campos obligatorios", a, len(df))

    a = len(df)
    df = df.drop_duplicates(subset=["url"], keep="first").copy()
    reg("eliminar_filas", "url", f"{a - len(df)} URL repetidas", "misma noticia recogida por varios temas", a, len(df))
    a = len(df)
    df["_clave"] = df["titulo"].map(norm_txt) + "|" + df["medio"].map(norm_txt)
    df = df.drop_duplicates(subset=["_clave"], keep="first").drop(columns="_clave").copy()
    reg("eliminar_filas", "titulo + medio", f"{a - len(df)} duplicados", "misma noticia con URL distinta", a, len(df))

    df["departamento_mencionado"] = [detectar_departamento(f"{t} {r}") for t, r in zip(df["titulo"], df["resumen"])]
    df["tiene_departamento"] = df["departamento_mencionado"] != ""
    reg("agregar_columna", "departamento_mencionado", f"{int((~df['tiene_departamento']).sum())} noticias sin departamento",
        "se marca en vez de eliminar; no se pueden cruzar por departamento", len(df), len(df))
    df["tema"] = df["tema_busqueda"].map(clasificar_tema)
    reg("agregar_columna", "tema", "tema normalizado desde tema_busqueda", "facilitar conteos en gold", len(df), len(df))

    df["archivo_origen"] = clave
    df["fecha_carga_silver"] = pd.Timestamp(datetime.now(ZoneInfo(TZ)).replace(tzinfo=None))

    s = df[["fecha", "hora_lima", "titulo", "medio", "url", "resumen", "tema", "tema_busqueda",
            "departamento_consulta", "departamento_mencionado", "tiene_departamento", "fecha_extraccion",
            "archivo_origen", "fecha_carga_silver"]].sort_values(["fecha", "hora_lima"]).reset_index(drop=True)

    errores = []
    if s["url"].duplicated().any():
        errores.append("hay URL duplicadas")
    for c in ["fecha", "titulo", "url"]:
        n = int(s[c].isna().sum())
        if n:
            errores.append(f"{c} tiene {n} nulos")
    if errores:
        raise ValueError("Validación fallida: " + "; ".join(errores))
    reg("validar", "url, fecha, titulo", "sin duplicados y sin nulos", "control de calidad antes de subir", len(s), len(s))

    despues = calidad_reporte(s.assign(fecha=s["fecha"].dt.strftime("%d-%m-%Y")).astype(str))

    stem = basename(clave)[:-4]
    carpeta = f"silver/{REPORTE}/{FECHA}"
    subir_bytes(f"{carpeta}/{stem}.xlsx",
                a_xlsx(s, {"fecha": "DD-MM-YYYY", "fecha_carga_silver": "DD-MM-YYYY HH:MM:SS"}), XLSX)
    subir_json(f"{carpeta}/{stem}_calidad.json", {
        "archivo_origen": clave, "fecha_proceso": FECHA,
        "antes": antes, "despues": despues, "reglas_aplicadas": reglas,
    })
    cols_silver = list(s.columns)
    subir_json(f"{carpeta}/{stem}.json", {
        "capa": "silver", "fuente": REPORTE_FUENTE, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "archivo_origen": clave, "fecha_origen": f_bronze,
        "archivo_destino": f"{carpeta}/{stem}.xlsx",
        "archivo_calidad": f"{carpeta}/{stem}_calidad.json",
        "periodo": f"{INICIO} a {FIN}",
        "filas_entrada": n0, "filas_salida": len(s),
        "columnas_agregadas": [c for c in cols_silver if c not in cols_bronze],
        "columna_final": cols_silver[-1],
        "tipos_de_datos": {c: str(t) for c, t in s.dtypes.items()},
        "estado": "procesado",
    })
    print(f"silver reporte: {n0:,} -> {len(s):,} filas -> s3://{BUCKET}/{carpeta}/{stem}.xlsx")


def reporte_gold():
    f_silver, clave = ultimo_simple("silver", REPORTE, ".xlsx")
    df = pd.read_excel(io.BytesIO(leer_bytes(clave)), sheet_name="datos")
    df["departamento_mencionado"] = df["departamento_mencionado"].fillna("").astype(str)
    for t in TEMAS_REPORTE:
        df[f"n_{t}"] = (df["tema"] == t).astype(int)

    agg = dict(n_noticias=("url", "count"), n_medios=("medio", "nunique"),
               **{f"n_{t}": (f"n_{t}", "sum") for t in TEMAS_REPORTE})

    # Para la integración: fecha + departamento (solo noticias con departamento detectado)
    dep = (df[df["departamento_mencionado"] != ""]
           .groupby(["fecha", "departamento_mencionado"], as_index=False).agg(**agg)
           .rename(columns={"departamento_mencionado": "departamento"})
           .sort_values(["fecha", "departamento"]))
    if dep.empty:
        raise ValueError("Ninguna noticia menciona un departamento: no hay nada que cruzar en la integración")
    # Visión nacional: todas las noticias por día
    nac = df.groupby("fecha", as_index=False).agg(**agg).sort_values("fecha")

    anios = df["fecha"].dt.year
    etiqueta = str(anios.min()) if anios.min() == anios.max() else f"{anios.min()}_{anios.max()}"
    dep["fecha"] = dep["fecha"].dt.strftime("%d-%m-%Y")
    nac["fecha"] = nac["fecha"].dt.strftime("%d-%m-%Y")

    carpeta = f"gold/{REPORTE}/{FECHA}"
    meta = []
    for nombre, frame in {"departamental": dep, "nacional": nac}.items():
        destino = f"{carpeta}/reporte_{nombre}_{etiqueta}_{FECHA}.csv"
        subir(destino, frame.to_csv(sep=SEP, index=False, na_rep=""))
        meta.append({"dataset": nombre, "registros": int(len(frame)),
                     "columnas": list(frame.columns), "ubicacion": f"s3://{BUCKET}/{destino}"})
        print(f"gold reporte/{nombre}: {len(frame):,} filas -> s3://{BUCKET}/{destino}")
    subir_json(f"{carpeta}/reporte_gold_{FECHA}.json", {
        "capa": "gold", "fuente": REPORTE_FUENTE, "fecha_proceso": FECHA, "zona_horaria": TZ,
        "origen_silver": clave, "periodo": f"{INICIO} a {FIN}",
        "dataset_para_integracion": "departamental (fecha + departamento)",
        "noticias_sin_departamento": int((df["departamento_mencionado"] == "").sum()),
        "datasets_gold": meta, "estado": "procesado",
    })


PASOS = {
    ("open_meteo", "fuente"): om_fuente, ("open_meteo", "bronze"): om_bronze,
    ("open_meteo", "silver"): om_silver, ("open_meteo", "gold"): om_gold,
    ("igp", "fuente"): igp_fuente, ("igp", "bronze"): igp_bronze,
    ("igp", "silver"): igp_silver, ("igp", "gold"): igp_gold,
    ("indeci", "fuente"): indeci_fuente, ("indeci", "bronze"): indeci_bronze,
    ("indeci", "silver"): indeci_silver, ("indeci", "gold"): indeci_gold,
    ("ingemmet", "fuente"): ingemmet_fuente, ("ingemmet", "bronze"): ingemmet_bronze,
    ("ingemmet", "silver"): ingemmet_silver, ("ingemmet", "gold"): ingemmet_gold,
    ("reporte", "fuente"): reporte_fuente, ("reporte", "bronze"): reporte_bronze,
    ("reporte", "silver"): reporte_silver, ("reporte", "gold"): reporte_gold,
    ("integracion", "gold"): gold_integracion,
}

renovar_credenciales()
print(f"Fecha de proceso (Lima): {FECHA}")

for fuente in FUENTES:
    for etapa in ETAPAS:
        paso = PASOS.get((fuente, etapa))
        if paso is None:
            continue
        print(f"\n=== {fuente.upper()} / {etapa.upper()} ===")
        try:
            paso()
        except Exception as e:
            print(f"{fuente}/{etapa} falló: {e!r}")
            break

print(f"\nTerminado ({datetime.now(ZoneInfo(TZ)).strftime('%Y-%m-%d %H:%M:%S')} Lima)")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Lectura de archivos (CSV / Excel) desde S3 hacia Spark
# MAGIC Usa la misma conexión del pipeline. Las columnas llegan como texto: castea las que necesites.

# COMMAND ----------

def leer_s3_csv_spark(clave, sep=SEP):
    pdf = pd.read_csv(io.BytesIO(leer_bytes(clave)), sep=sep, dtype=str, keep_default_na=False)
    return spark.createDataFrame(pdf)


def leer_s3_excel_spark(clave, hoja="datos"):
    pdf = pd.read_excel(io.BytesIO(leer_bytes(clave)), sheet_name=hoja, dtype=str).fillna("")
    return spark.createDataFrame(pdf)


f_ult, clave_csv = ultimo_simple("gold", "integracion_de_datos", ".csv")
df_integracion = leer_s3_csv_spark(clave_csv)
print(f"s3://{BUCKET}/{clave_csv}: {df_integracion.count():,} filas, {len(df_integracion.columns)} columnas")
display(df_integracion.limit(10))

f_ult, clave_xlsx = ultimo_simple("gold", "integracion_de_datos", ".xlsx")
df_excel = leer_s3_excel_spark(clave_xlsx, hoja=pd.ExcelFile(io.BytesIO(leer_bytes(clave_xlsx))).sheet_names[0])
print(f"s3://{BUCKET}/{clave_xlsx}: {df_excel.count():,} filas (Excel)")